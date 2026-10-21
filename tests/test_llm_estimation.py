"""decisions #19：缺烹調換算（產率／保留率）與容量單位克數時，改用事先批次推估並快取的 LLM 係數。

測試的介面（seams）：
1. make_tfda_calculator(cooking_factors=...)：純計算，真實 T04 TFDA 資料 + 快取係數 → 營養（Decimal）與追溯標示；
2. load_recipe_lines(db, ...)：容量單位克數推估解出每 1n 克數；
3. estimation.store：查表優先順序（TFDA 實測 > 人工 > LLM；rejected 不用）與覆核；
4. estimation.batch：掃描缺口、dry-run、冪等、Structured Outputs 驗證、T10 預留／結算、無金鑰錯誤；
5. 結果檢視：受影響食材與營養標示「LLM 推估」。
絕不呼叫真實 OpenAI：responder 一律為假物件。
"""
import json
import tempfile
import unittest
import uuid as uuid_module
from decimal import Decimal
from fractions import Fraction
from pathlib import Path

from app.estimation import batch, store
from app.meal_calc.types import EstimatedFactor
from app.services import meal_calculation as mc
from db_support import migrate_fresh, requires_db, session_factory, truncate_all
from recipe_workbook import build, ingredient, recipe

D = Decimal


def line(name, grams, food_id=None, **kw):
    return mc.IngredientLine(food_id=food_id or f"F-{name}", name=name,
                             grams_per_n=None if grams is None else D(grams), **kw)


def cooking(food_id, method, *, yield_factor="0.75", energy="0.95", protein="0.9", status="llm_estimated",
            source="llm_estimated", estimate_id="est-1"):
    return EstimatedFactor(
        estimate_id=estimate_id, kind="cooking", food_id=food_id, key=method, source=source, status=status,
        model="gpt-5.4-mini-2026-03-17", prompt_version="p1", schema_version="s1",
        yield_factor=D(yield_factor) if yield_factor is not None else None,
        retention={"energy_kcal": D(energy), "protein_g": D(protein)})


def lookup(*factors):
    table = {(f.food_id, f.key): f for f in factors}
    return lambda food_id, method: table.get((food_id, mc.normalize_cooking_method(method)))


class CachedCookingFactorTests(unittest.TestCase):
    """真實 T04 TFDA 資料；係數來自快取，計算期間不呼叫任何 LLM。"""

    def test_raw_food_with_cached_retention_becomes_complete_and_is_marked_llm_estimated(self):
        # 雞胸肉 I0402402 生：117 kcal、23.3 g 蛋白質／100 g；150 g × 保留率 0.95／0.9
        calc = mc.make_tfda_calculator(cooking_factors=lookup(cooking("F-雞胸肉", "roast")))
        [chicken] = calc([(line("雞胸肉", "150", cooking_method="烤"), D("150"))])
        self.assertIsNone(chicken.error)
        self.assertEqual(chicken.nutrients["energy_kcal"], D("166.725"))   # 117 × 1.5 × 0.95
        self.assertEqual(chicken.nutrients["protein_g"], D("31.455"))      # 23.3 × 1.5 × 0.9
        # 非必要營養素沒有推估保留率：decisions #40 以生組成、保留率 1 近似；TFDA 空白者仍未知、不填 0
        self.assertEqual(chicken.nutrients["fat_g"], D("2.85"))
        self.assertIsNone(chicken.nutrients["dietary_fiber_g"])
        self.assertTrue(chicken.trace["is_estimated"])
        self.assertEqual(chicken.trace["sample_id"], "I0402402")
        [marker] = chicken.trace["estimated_factors"]
        self.assertEqual((marker["kind"], marker["source"], marker["status"], marker["model"],
                          marker["prompt_version"], marker["schema_version"]),
                         ("cooking", "llm_estimated", "llm_estimated", "gpt-5.4-mini-2026-03-17", "p1", "s1"))
        self.assertEqual(marker["retention"], {"energy_kcal": D("0.95"), "protein_g": D("0.9")})
        self.assertEqual(sorted(chicken.trace["estimated_nutrients"]), ["energy_kcal", "protein_g"])

    def test_cooked_tfda_record_uses_cached_yield_for_cooked_weight(self):
        # 白芝麻 烤 → TFDA 熟樣品 C1610101（626 kcal、20.3 g／100 g），無 USDA 產率；10 g 生 × 0.95 = 9.5 g 熟
        calc = mc.make_tfda_calculator(cooking_factors=lookup(cooking("F-白芝麻", "roast", yield_factor="0.95")))
        [sesame] = calc([(line("白芝麻", "10", cooking_method="烤"), D("10"))])
        self.assertIsNone(sesame.error)
        self.assertEqual(sesame.nutrients["energy_kcal"], D("59.47"))
        self.assertEqual(sesame.nutrients["protein_g"], D("1.9285"))
        self.assertEqual(sesame.trace["cooked_weight_g"], D("9.5"))
        self.assertEqual(sesame.trace["sample_id"], "C1610101")
        self.assertEqual(sesame.trace["estimated_factors"][0]["yield_factor"], D("0.95"))

    def test_without_a_cached_factor_the_old_behaviour_stays(self):
        calc = mc.make_tfda_calculator(cooking_factors=lookup())
        [chicken] = calc([(line("雞胸肉", "150", cooking_method="烤"), D("150"))])
        self.assertIsNone(chicken.nutrients["energy_kcal"])
        self.assertFalse(chicken.trace["is_estimated"])
        self.assertEqual(chicken.trace.get("estimated_factors"), [])

    def test_rejected_factor_is_never_used(self):
        calc = mc.make_tfda_calculator(cooking_factors=lookup(cooking("F-雞胸肉", "roast", status="rejected")))
        [chicken] = calc([(line("雞胸肉", "150", cooking_method="烤"), D("150"))])
        self.assertIsNone(chicken.nutrients["energy_kcal"])

    def test_tfda_measured_data_wins_over_any_cached_factor(self):
        # 白飯為 TFDA 熟樣品、以熟重計：不套任何推估
        calc = mc.make_tfda_calculator(cooking_factors=lookup(cooking("F-白飯", "steam", yield_factor="2")))
        [rice] = calc([(line("白飯", "200", weight_basis="cooked", cooking_method="蒸"), D("200"))])
        self.assertEqual(rice.nutrients["energy_kcal"], D("366"))
        self.assertEqual(rice.trace["estimated_factors"], [])

    def test_meal_using_an_estimate_is_adoptable_and_reports_which_values_are_estimated(self):
        calc = mc.make_tfda_calculator(cooking_factors=lookup(cooking("F-雞胸肉", "roast")))
        # 150 g 雞胸肉 → 166.725 kcal／31.455 g；目標剛好以此為中心
        targets = mc.MealTargets(Fraction(1), Fraction("166.725"), Fraction("31.455"))
        result = mc.calculate_meal([line("雞胸肉", "150", cooking_method="烤")], targets, mode="ai",
                                   calculator=calc)
        self.assertTrue(result["adoptable"], result["block_reasons"])
        [row] = result["ingredients"]
        self.assertTrue(row["llm_estimated"])
        self.assertEqual(result["llm_estimates"]["food_ids"], ["F-雞胸肉"])
        self.assertEqual(sorted(result["llm_estimates"]["nutrients"]), ["energy_kcal", "protein_g"])
        self.assertTrue(result["llm_estimates"]["used"])

    def test_recipe_with_several_methods_uses_the_combined_method_key(self):
        # 多種烹調方式的食譜：T12 不猜是哪一種，改以「全部方式」組合為鍵查推估（批次同樣以此鍵推估）
        key = mc.combined_method_key(["炒", "煮"])
        self.assertEqual(key, "+".join(sorted([mc.normalize_cooking_method("炒"), mc.normalize_cooking_method("煮")])))
        calc = mc.make_tfda_calculator(cooking_factors=lookup(cooking("F-雞胸肉", key)))
        ln = line("雞胸肉", "150", cooking_methods=("炒", "煮"))
        self.assertIsNone(ln.cooking_method)
        [chicken] = calc([(ln, D("150"))])
        self.assertEqual(chicken.nutrients["energy_kcal"], D("166.725"))
        self.assertEqual(chicken.trace["estimated_factors"][0]["key"], key)
        [plain] = mc.tfda_calculator([(ln, D("150"))])
        self.assertIsNone(plain.nutrients["energy_kcal"])  # 沒有推估值：維持不可採用

    def test_default_tfda_calculator_never_calls_a_configured_llm(self):
        import os
        from unittest.mock import patch
        with patch.dict(os.environ, {"GEMINI_API_KEY": "should-not-be-used", "NUTRITION_LLM_PROVIDER": "gemini"}):
            sys_mod = mc._nutrition_model()
            import llm_cooking_estimator as est
            with patch.object(est, "urlopen", side_effect=AssertionError("LLM called at request time")):
                [chicken] = mc.tfda_calculator([(line("雞胸肉", "150", cooking_method="烤"), D("150"))])
        self.assertIsNone(chicken.nutrients["energy_kcal"])
        self.assertIsNotNone(sys_mod)


MODEL = "gpt-5.4-mini-2026-03-17"


@requires_db
class FactorStoreTests(unittest.TestCase):
    """快取表：TFDA 實測（不入表）> 人工核可／覆寫 > LLM 推估；rejected 永不使用。"""

    @classmethod
    def setUpClass(cls):
        migrate_fresh()

    def setUp(self):
        truncate_all()

    def add_llm(self, db, food_id="FOOD000100", method="roast", energy="0.95"):
        return store.add_estimate(
            db, kind="cooking", food_id=food_id, key=method, food_name="雞胸肉", yield_factor=D("0.75"),
            retention={"energy_kcal": D(energy), "protein_g": D("0.9")}, source="llm_estimated",
            status="llm_estimated", model=MODEL, prompt_version="p1", schema_version="s1",
            input_snapshot={"food_name": "雞胸肉"}, assumptions=["去皮烤"], call_key="k1")

    def test_llm_estimate_is_looked_up_with_exact_decimal_values(self):
        with session_factory()() as db, db.begin():
            self.add_llm(db)
        with session_factory()() as db:
            snap = store.load_snapshot(db)
        factor = snap.cooking("FOOD000100", "roast")
        self.assertEqual((factor.source, factor.status, factor.model), ("llm_estimated", "llm_estimated", MODEL))
        self.assertEqual(factor.retention, {"energy_kcal": D("0.95"), "protein_g": D("0.9")})
        self.assertEqual(factor.yield_factor, D("0.75"))
        self.assertIsNone(snap.cooking("FOOD000100", "boil"))

    def test_human_override_wins_and_rejecting_it_falls_back_to_llm(self):
        with session_factory()() as db, db.begin():
            llm = self.add_llm(db)
            human = store.override(db, kind="cooking", food_id="FOOD000100", key="roast", yield_factor=D("0.7"),
                                   retention={"energy_kcal": D("1"), "protein_g": D("1")},
                                   reviewer="營養師A", note="依實測")
        with session_factory()() as db:
            factor = store.load_snapshot(db).cooking("FOOD000100", "roast")
        self.assertEqual((factor.estimate_id, factor.source, factor.status), (str(human), "human", "overridden"))
        self.assertFalse(factor.llm_estimated)
        with session_factory()() as db, db.begin():
            store.review(db, human, action="reject", reviewer="營養師B", note="誤植")
        with session_factory()() as db:
            factor = store.load_snapshot(db).cooking("FOOD000100", "roast")
        self.assertEqual(factor.estimate_id, str(llm))

    def test_approved_llm_estimate_beats_a_newer_unreviewed_one(self):
        with session_factory()() as db, db.begin():
            old = self.add_llm(db, energy="0.9")
            store.review(db, old, action="approve", reviewer="營養師A", note=None)
            self.add_llm(db, energy="0.8")
        with session_factory()() as db:
            factor = store.load_snapshot(db).cooking("FOOD000100", "roast")
        self.assertEqual((factor.estimate_id, factor.status), (str(old), "approved"))
        self.assertTrue(factor.llm_estimated)  # 數值仍來自 LLM，結果照樣標示

    def test_rejected_only_means_no_factor(self):
        with session_factory()() as db, db.begin():
            store.review(db, self.add_llm(db), action="reject", reviewer="營養師A", note="不合理")
        with session_factory()() as db:
            self.assertIsNone(store.load_snapshot(db).cooking("FOOD000100", "roast"))

    def test_unit_estimates_and_review_listing(self):
        with session_factory()() as db, db.begin():
            store.add_estimate(db, kind="unit", food_id="FOOD001541", key="匙", food_name="醬油",
                               grams_per_unit=D("15.5"), source="llm_estimated", status="llm_estimated",
                               model=MODEL, prompt_version="p1", schema_version="s1",
                               input_snapshot={"unit": "匙"}, assumptions=["一匙約一大匙"], call_key="k2")
        with session_factory()() as db:
            factor = store.load_snapshot(db).unit("FOOD001541", "匙")
            rows = store.list_estimates(db, status="llm_estimated")
        self.assertEqual(factor.grams_per_unit, D("15.5"))
        [row] = rows
        self.assertEqual((row["kind"], row["food_id"], row["key"], row["grams_per_unit"], row["created_at"] is not None),
                         ("unit", "FOOD001541", "匙", D("15.5"), True))

    def test_values_outside_plausible_ranges_are_refused_by_the_database(self):
        from sqlalchemy.exc import IntegrityError
        with self.assertRaises(IntegrityError):
            with session_factory()() as db, db.begin():
                self.add_llm(db, energy="1.5")


@requires_db
class CachedFactorsInRecipeCalculationTests(unittest.TestCase):
    """T12 載入食譜與計算時使用快取（不呼叫 LLM）。"""

    @classmethod
    def setUpClass(cls):
        migrate_fresh()

    def setUp(self):
        truncate_all()
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        from app.services.recipe_import import import_recipe_bundle, parse_recipe_workbook
        bundle = parse_recipe_workbook(build(Path(self.folder.name) / "book.xlsx", [
            recipe("R001", "烤雞胸", servings="2人份", ingredients=[
                ingredient("FOOD000100", "雞胸肉", 300),
                ingredient("FOOD001541", "醬油", 1, raw_text="醬油 1匙", unit="匙", weight_g=None),
                ingredient("FOOD000039", "大蒜", None, raw_text="大蒜 適量", unit=None, amount=None, weight_g=None),
            ])]))
        with session_factory()() as session, session.begin():
            import_recipe_bundle(session, bundle)

    def add_unit(self, grams="15", status="llm_estimated"):
        with session_factory()() as db, db.begin():
            return store.add_estimate(db, kind="unit", food_id="FOOD001541", key="匙", food_name="醬油",
                                      grams_per_unit=D(grams), source="llm_estimated", status=status, model=MODEL,
                                      prompt_version="p1", schema_version="s1", input_snapshot={}, call_key="k")

    def load(self):
        with session_factory()() as db:
            return mc.load_recipe_lines(db, "R001")

    def test_unit_estimate_resolves_grams_per_n_and_is_traceable(self):
        estimate_id = self.add_unit("15")
        _, soy, garlic = self.load()
        self.assertEqual(soy.grams_per_n, D("7.5"))  # 1 匙 × 15 g ÷ 2 人份
        self.assertIsNone(soy.unresolved_reason)
        self.assertEqual((soy.weight_estimate.estimate_id, soy.weight_estimate.source), (str(estimate_id),
                                                                                         "llm_estimated"))
        self.assertIsNone(garlic.grams_per_n)  # 「適量」沒有數量：推估單位克數也無法換算
        result = mc.calculate_meal([soy], mc.MealTargets(Fraction(1), Fraction(1), Fraction(1)), mode="ai",
                                   calculator=mc.tfda_calculator)
        self.assertNotIn("INGREDIENT_UNRESOLVED", [r["code"] for r in result["block_reasons"]])
        [row] = result["ingredients"]
        self.assertTrue(row["llm_estimated"])
        self.assertEqual(row["estimates"][0]["kind"], "unit")
        self.assertEqual(row["whole_meal_grams"], D("7.5"))

    def import_more(self, *recipes):
        from app.services.recipe_import import import_recipe_bundle, parse_recipe_workbook
        bundle = parse_recipe_workbook(build(Path(self.folder.name) / "more.xlsx", list(recipes)))
        with session_factory()() as session, session.begin():
            import_recipe_bundle(session, bundle)

    def test_workbook_weight_wins_over_unit_estimate_and_is_not_labelled_llm(self):
        # decisions #20：該列有工作簿 weight_g 就直接採用，即使已有 LLM 單位推估
        self.add_unit("15")
        self.import_more(recipe("R009", "醬燒雞胸", servings="2人份", ingredients=[
            ingredient("FOOD001541", "醬油", 1, raw_text="醬油 1大匙", unit="大匙", weight_g=18)]))
        with session_factory()() as db:
            [soy] = mc.load_recipe_lines(db, "R009")
        self.assertEqual(soy.grams_per_n, D("9"))  # 18 g ÷ 2 人份
        self.assertIsNone(soy.weight_estimate)
        self.assertEqual(soy.weight_source, "workbook")
        result = mc.calculate_meal([soy], mc.MealTargets(Fraction(1), Fraction(1), Fraction(1)), mode="ai",
                                   calculator=mc.tfda_calculator)
        [row] = result["ingredients"]
        self.assertFalse(row["llm_estimated"])
        self.assertEqual(row["estimates"], [])
        self.assertEqual(row["weight_source"], "workbook")
        self.assertEqual(row["whole_meal_grams"], D("9"))

    def test_weight_source_marks_gram_lines_and_unit_estimates(self):
        self.add_unit("15")
        chicken, soy, garlic = self.load()
        self.assertEqual((chicken.weight_source, soy.weight_source, garlic.weight_source),
                         ("recipe_grams", "unit_estimate", None))

    def test_without_unit_estimate_the_line_stays_unresolved(self):
        _, soy, _ = self.load()
        self.assertIsNone(soy.grams_per_n)
        self.assertIsNone(soy.weight_estimate)

    def test_rejected_unit_estimate_is_not_used(self):
        self.add_unit("15", status="rejected")
        _, soy, _ = self.load()
        self.assertIsNone(soy.grams_per_n)

    def test_calculator_from_db_applies_cached_cooking_factor(self):
        with session_factory()() as db, db.begin():
            store.add_estimate(db, kind="cooking", food_id="FOOD000100", key="炒", food_name="雞胸肉",
                               yield_factor=D("0.8"), retention={"energy_kcal": D("1"), "protein_g": D("0.9")},
                               source="llm_estimated", status="llm_estimated", model=MODEL, prompt_version="p1",
                               schema_version="s1", input_snapshot={}, call_key="k")
        chicken, _, _ = self.load()
        with session_factory()() as db:
            calc = mc.cached_factor_calculator(db)
        [row] = calc([(chicken, chicken.grams_per_n)])
        self.assertEqual(row.nutrients["energy_kcal"], D("175.5"))  # 117 × 1.5 × 1
        self.assertEqual(row.trace["estimated_factors"][0]["key"], mc.normalize_cooking_method("炒"))


class FakeResponder:
    """假的 OpenAI Responses 轉接器：依 schema_name 回傳預設內容，記錄每次呼叫。"""

    def __init__(self, cooking=None, unit=None):
        self.cooking = cooking if cooking is not None else {
            "yield_factor": 0.75, "energy_retention": 0.95, "protein_retention": 0.9,
            "assumptions": ["去皮、未瀝湯汁"]}
        self.unit = unit if unit is not None else {"grams_per_unit": 15, "assumptions": ["一匙視為一大匙"]}
        self.calls = []

    def respond(self, **kwargs):
        from app.services.llm_client import StructuredResponse
        self.calls.append(kwargs)
        outcome = self.cooking if kwargs["schema_name"] == "cooking_factor_estimate" else self.unit
        if isinstance(outcome, Exception):
            raise outcome
        if isinstance(outcome, StructuredResponse):
            return outcome, {"input_tokens": 400, "output_tokens": 80}
        text = outcome if isinstance(outcome, str) else json.dumps(outcome, ensure_ascii=False)
        return StructuredResponse("completed", output_text=text), {"input_tokens": 400, "output_tokens": 80}


@requires_db
class BatchEstimationTests(unittest.TestCase):
    """批次指令：掃描缺口 → gpt-5.4-mini 推估 → 驗證 → 寫入快取；每次呼叫經 T10 預留／結算。"""

    @classmethod
    def setUpClass(cls):
        migrate_fresh()

    def setUp(self):
        truncate_all()
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        from app.services.recipe_import import import_recipe_bundle, parse_recipe_workbook
        bundle = parse_recipe_workbook(build(Path(self.folder.name) / "book.xlsx", [
            recipe("R001", "炒雞胸", servings="2人份", ingredients=[
                ingredient("FOOD000100", "雞胸肉", 300),
                ingredient("FOOD001541", "醬油", 1, raw_text="醬油 1匙", unit="匙", weight_g=None),
            ]),
            recipe("R002", "雞胸沙拉", servings="1人份", ingredients=[ingredient("FOOD000100", "雞胸肉", 100)]),
        ]))
        with session_factory()() as session, session.begin():
            import_recipe_bundle(session, bundle)

    def reservations(self):
        from sqlalchemy import select
        from app.db.models import UsageReservation
        with session_factory()() as db:
            return db.scalars(select(UsageReservation)).all()

    def rows(self, **kw):
        with session_factory()() as db:
            return store.list_estimates(db, **kw)

    def run_batch(self, responder, **kw):
        return batch.run(session_factory(), responder=responder, **kw)

    def test_scan_finds_missing_cooking_and_unit_pairs_once_each(self):
        with session_factory()() as db:
            gaps = batch.scan(db)
        keys = {(g.kind, g.food_id, g.key) for g in gaps}
        stir = mc.normalize_cooking_method("炒")
        self.assertIn(("cooking", "FOOD000100", stir), keys)
        self.assertIn(("unit", "FOOD001541", "匙"), keys)
        self.assertEqual(len([g for g in gaps if (g.kind, g.food_id) == ("cooking", "FOOD000100")]), 1)
        chicken = next(g for g in gaps if (g.kind, g.food_id) == ("cooking", "FOOD000100"))
        self.assertEqual(chicken.input_snapshot["tfda_raw"]["sample_id"], "I0402402")
        self.assertEqual(sorted(chicken.input_snapshot["recipe_ids"]), ["R001", "R002"])

    def test_lines_with_workbook_weight_never_become_unit_gaps(self):
        # decisions #20：以真實工作簿的形狀（每種非 g 單位各取一列，weight_g 皆有值）匯入，單位缺口為 0
        from openpyxl import load_workbook
        from app.core.config import PROJECT_ROOT
        rows = load_workbook(PROJECT_ROOT / "data/recipe_dataset_1017.xlsx", read_only=True,
                             data_only=True)["recipe_ingredients"].iter_rows(values_only=True)
        header = list(next(rows))
        col = {name: header.index(name) for name in ("raw_text", "amount", "unit", "weight_g")}
        samples = {}
        for row in rows:
            unit = row[col["unit"]]
            if unit and unit != "g" and row[col["amount"]] is not None and unit not in samples:
                samples[unit] = row
        self.assertGreaterEqual(len(samples), 10)
        self.assertTrue(all(r[col["weight_g"]] is not None for r in samples.values()))
        lines = [ingredient("FOOD001541", "醬油", r[col["amount"]], raw_text=f"醬油 {r[col['amount']]}{u}",
                            unit=u, weight_g=r[col["weight_g"]]) for u, r in samples.items()]
        truncate_all()
        from app.services.recipe_import import import_recipe_bundle, parse_recipe_workbook
        bundle = parse_recipe_workbook(build(Path(self.folder.name) / "real.xlsx", [
            recipe("R100", "醬油試算", servings="2人份", ingredients=lines)]))
        with session_factory()() as session, session.begin():
            import_recipe_bundle(session, bundle)
        with session_factory()() as db:
            gaps = batch.scan(db)
            loaded = mc.load_recipe_lines(db, "R100")
        self.assertEqual([g for g in gaps if g.kind == "unit"], [])
        resolved = [ln for ln in loaded if ln.unit and ln.unit != "g"]
        self.assertTrue(resolved)
        self.assertTrue(all(ln.grams_per_n is not None and ln.weight_source == "workbook" for ln in resolved))

    def test_recipe_with_several_methods_is_scanned_under_the_combined_key_and_then_adoptable(self):
        from sqlalchemy import text
        with session_factory()() as db, db.begin():
            db.execute(text("""UPDATE recipes SET cooking_methods = '["炒", "煮"]' WHERE recipe_id = 'R002'"""))
        with session_factory()() as db:
            keys = {(g.kind, g.food_id, g.key) for g in batch.scan(db)}
        combo = mc.combined_method_key(["炒", "煮"])
        self.assertIn(("cooking", "FOOD000100", combo), keys)
        self.run_batch(FakeResponder(cooking={"yield_factor": 0.8, "energy_retention": 1, "protein_retention": 1,
                                              "assumptions": ["x"]}))
        with session_factory()() as db:
            lines = mc.load_recipe_lines(db, "R002")
            calc = mc.cached_factor_calculator(db)
        result = mc.calculate_meal(lines, mc.MealTargets(Fraction(1), Fraction(117), Fraction("23.3")), mode="ai",
                                   calculator=calc)
        self.assertTrue(result["adoptable"], result["block_reasons"])

    def test_dry_run_lists_pairs_without_calling_the_model_or_writing(self):
        responder = FakeResponder()
        report = self.run_batch(responder, dry_run=True)
        self.assertTrue(report["dry_run"])
        self.assertGreaterEqual(report["pending"], 2)
        self.assertEqual(responder.calls, [])
        self.assertEqual(self.rows(), [])
        self.assertEqual(self.reservations(), [])

    def test_estimates_are_validated_stored_with_trace_and_billed_through_t10(self):
        responder = FakeResponder()
        report = self.run_batch(responder)
        self.assertEqual(report["failed"], [])
        self.assertEqual(report["estimated"], len(responder.calls))
        call = responder.calls[0]
        self.assertEqual(call["model"], MODEL)
        self.assertTrue(call["schema"]["additionalProperties"] is False)
        [chicken] = self.rows(kind="cooking", food_id="FOOD000100")
        self.assertEqual((chicken["source"], chicken["status"], chicken["model"]), ("llm_estimated", "llm_estimated", MODEL))
        self.assertEqual((chicken["prompt_version"], chicken["schema_version"]),
                         (batch.COOKING_PROMPT_VERSION, batch.COOKING_SCHEMA_VERSION))
        self.assertEqual((chicken["yield_factor"], chicken["energy_retention"], chicken["protein_retention"]),
                         (D("0.75"), D("0.95"), D("0.9")))
        [soy] = self.rows(kind="unit", food_id="FOOD001541")
        self.assertEqual(soy["grams_per_unit"], D("15"))
        reservations = self.reservations()
        self.assertEqual(len(reservations), len(responder.calls))
        self.assertTrue(all(r.status == "settled" and r.purpose == "conversion_factor_estimation"
                            for r in reservations))

    def test_rerun_is_idempotent_unless_refresh(self):
        self.run_batch(FakeResponder())
        again = FakeResponder()
        report = self.run_batch(again)
        self.assertEqual(again.calls, [])
        self.assertEqual(report["pending"], 0)
        self.assertGreater(report["skipped_existing"], 0)
        refresh = FakeResponder()
        self.run_batch(refresh, refresh=True, limit=1)
        self.assertEqual(len(refresh.calls), 1)

    def test_refusal_truncation_bad_format_and_out_of_range_are_rejected_and_not_written(self):
        from app.services.llm_client import StructuredResponse
        cases = {
            "MODEL_REFUSED": StructuredResponse("refused", refusal="no"),
            "OUTPUT_TRUNCATED": StructuredResponse("incomplete", incomplete_reason="max_output_tokens"),
            "INVALID_FORMAT": "{not json",
            "OUT_OF_RANGE": {"yield_factor": 0.75, "energy_retention": 0.2, "protein_retention": 0.9,
                             "assumptions": ["x"]},
        }
        for reason, outcome in cases.items():
            with self.subTest(reason):
                truncate_all_estimates()
                report = self.run_batch(FakeResponder(cooking=outcome, unit={"grams_per_unit": 500,
                                                                             "assumptions": ["x"]}))
                self.assertIn(reason, {f["reason"] for f in report["failed"] if f["kind"] == "cooking"})
                self.assertIn("OUT_OF_RANGE", {f["reason"] for f in report["failed"] if f["kind"] == "unit"})
                self.assertEqual(self.rows(), [])

    def test_without_api_key_it_fails_clearly_and_writes_nothing(self):
        with self.assertRaises(batch.EstimationUnavailable) as ctx:
            self.run_batch(None)
        self.assertIn("OPENAI_API_KEY", str(ctx.exception))
        self.assertEqual(self.rows(), [])
        self.assertEqual(self.reservations(), [])

    def test_budget_exhausted_stops_the_batch(self):
        from sqlalchemy import text
        from app.services import spend
        with session_factory()() as db, db.begin():
            spend.ensure_month(db, spend.month_key())
            db.execute(text("UPDATE budget_months SET spent_twd = limit_twd - committed_twd"))
        responder = FakeResponder()
        report = self.run_batch(responder)
        self.assertEqual(report["stopped"], "BUDGET_EXHAUSTED")
        self.assertEqual(responder.calls, [])
        self.assertEqual(self.rows(), [])

    def test_estimated_recipe_becomes_adoptable_end_to_end(self):
        self.run_batch(FakeResponder(cooking={"yield_factor": 0.8, "energy_retention": 1, "protein_retention": 1,
                                              "assumptions": ["x"]}))
        with session_factory()() as db:
            lines = mc.load_recipe_lines(db, "R002")
            calc = mc.cached_factor_calculator(db)
        # 雞胸肉 100 g 生 × 保留率 1 → 117 kcal、23.3 g
        targets = mc.MealTargets(Fraction(1), Fraction(117), Fraction("23.3"))
        result = mc.calculate_meal(lines, targets, mode="ai", calculator=calc)
        self.assertTrue(result["adoptable"], result["block_reasons"])
        self.assertTrue(result["llm_estimates"]["used"])


@requires_db
class CommandLineTests(unittest.TestCase):
    """python -m app.estimation：run／list／approve／reject／override。"""

    @classmethod
    def setUpClass(cls):
        migrate_fresh()

    def setUp(self):
        truncate_all()

    def cli(self, *argv, responder=None):
        import io
        from contextlib import redirect_stderr, redirect_stdout
        from app.estimation import __main__ as cli
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = cli.main(list(argv), factory=session_factory(), responder=responder)
        return code, out.getvalue(), err.getvalue()

    def test_run_without_key_exits_with_clear_error(self):
        from unittest.mock import patch
        from app.core.config import settings
        with patch.object(settings, "openai_api_key", None):
            code, _, err = self.cli("run")
        self.assertEqual(code, 2)
        self.assertIn("OPENAI_API_KEY", err)

    def test_dry_run_prints_the_pending_pairs_as_json(self):
        code, out, _ = self.cli("run", "--dry-run")
        self.assertEqual(code, 0)
        self.assertTrue(json.loads(out)["dry_run"])

    def test_review_commands(self):
        with session_factory()() as db, db.begin():
            estimate_id = store.add_estimate(
                db, kind="unit", food_id="FOOD001541", key="匙", food_name="醬油", grams_per_unit=D("15"),
                source="llm_estimated", status="llm_estimated", model=MODEL, prompt_version="p1",
                schema_version="s1", input_snapshot={}, call_key="k")
        code, out, _ = self.cli("list", "--status", "llm_estimated")
        self.assertEqual((code, json.loads(out)[0]["id"]), (0, str(estimate_id)))
        self.assertEqual(self.cli("approve", str(estimate_id), "--reviewer", "營養師A")[0], 0)
        code, _, _ = self.cli("override", "--kind", "unit", "--food", "FOOD001541", "--key", "匙",
                              "--grams", "16.5", "--reviewer", "營養師A", "--note", "量測")
        self.assertEqual(code, 0)
        with session_factory()() as db:
            factor = store.load_snapshot(db).unit("FOOD001541", "匙")
        self.assertEqual((factor.source, factor.grams_per_unit), ("human", D("16.5")))
        code, _, err = self.cli("reject", str(uuid_module.uuid4()), "--reviewer", "營養師A")
        self.assertEqual(code, 1)
        self.assertIn("not found", err)


def truncate_all_estimates():
    from sqlalchemy import text
    with session_factory()() as db, db.begin():
        db.execute(text("DELETE FROM conversion_factor_estimates"))
        db.execute(text("DELETE FROM usage_reservations"))
        db.execute(text("UPDATE budget_months SET reserved_twd = 0, spent_twd = 0"))


class ResultViewLabelTests(unittest.TestCase):
    """T15 結果檢視：受影響食材與營養帶「LLM 推估」標示（前端依此顯示）。"""

    def test_affected_ingredients_and_nutrients_are_labelled(self):
        from app.services import meal_results
        calc = mc.make_tfda_calculator(cooking_factors=lookup(cooking("F-雞胸肉", "roast")))
        targets = mc.MealTargets(Fraction(1), Fraction("200"), Fraction("35"))
        result = mc.calculate_meal([line("雞胸肉", "150", cooking_method="烤", dish_id="R1"),
                                    line("大豆油", "5", consumed_as_is=True, dish_id="R1")], targets, mode="ai",
                                   calculator=calc)
        payload = {"calculation": result, "dishes": [{"standard_dish_id": "D1", "dish_name": "烤雞胸",
                                                      "source": "database", "recipe_id": "R1"}]}
        [dish] = meal_results._dishes(payload)
        chicken, oil = dish["ingredients"]
        self.assertTrue(chicken["llm_estimated"])
        self.assertEqual(chicken["estimate_label"], "LLM 推估")
        self.assertEqual(chicken["estimates"][0]["model"], MODEL)
        self.assertFalse(oil["llm_estimated"])
        self.assertIsNone(oil["estimate_label"])
        nutrition = meal_results._nutrition(payload)
        self.assertEqual(nutrition["llm_estimates"]["food_ids"], ["F-雞胸肉"])
        self.assertEqual(sorted(nutrition["llm_estimates"]["nutrients"]), ["energy_kcal", "protein_g"])
        self.assertEqual(nutrition["llm_estimates"]["label"], "LLM 推估")

    def test_old_payload_without_estimates_still_renders(self):
        from app.services import meal_results
        payload = {"calculation": {"ingredients": [{"food_id": "F1", "name": "x", "dish_id": "R1"}]},
                   "dishes": [{"standard_dish_id": "D1", "dish_name": "d", "source": "database", "recipe_id": "R1"}]}
        [dish] = meal_results._dishes(payload)
        self.assertFalse(dish["ingredients"][0]["llm_estimated"])
        self.assertFalse(meal_results._nutrition(payload)["llm_estimates"]["used"])


if __name__ == "__main__":
    unittest.main()
