"""T09：照片辨識到食材種類確認（AC01、AC04）。推論結果皆為假服務回傳的合成資料。"""
import io
import math
import unittest
from unittest.mock import patch

import httpx
from PIL import Image

from db_support import seed_foods, make_client, migrate_fresh, requires_db, truncate_all
from app.services.team_link import TeamCredentials

PASSWORD = "correct-horse-42"


def jpeg_bytes(size=(64, 32), *, orientation=None, gps=False) -> bytes:
    img = Image.new("RGB", size, (0, 160, 0))
    exif = Image.Exif()
    if orientation:
        exif[0x0112] = orientation
    if gps:
        exif[0x8825] = {1: "N", 2: (25.0, 2.0, 0.0)}
    out = io.BytesIO()
    img.save(out, format="JPEG", exif=exif.tobytes())
    return out.getvalue()


def png_bytes() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (20, 20), (200, 0, 0)).save(out, format="PNG")
    return out.getvalue()


class FakeRecognizer:
    def __init__(self, detections=None, *, error=None, version="yolov8n-synthetic-test"):
        self.detections = detections or []
        self.error = error
        self.version = version
        self.calls = 0

    def detect(self, jpeg):
        from app.services.recognition import Detection, RecognizerOutput
        self.calls += 1
        Image.open(io.BytesIO(jpeg)).verify()  # 送進推論的是已解碼重編的 JPEG
        if self.error:
            raise self.error
        return RecognizerOutput(self.version, [Detection(l, c, [0, 0, 10, 10]) for l, c in self.detections])


class SelectionTests(unittest.TestCase):
    def test_highest_valid_confidence_wins(self):
        from app.services.recognition import Detection, top_detection
        dets = [Detection("高麗菜", 0.62, None), Detection("青江菜", 0.91, None), Detection("", 0.99, None),
                Detection("番茄", float("nan"), None), Detection("洋蔥", 1.5, None), Detection("蔥", 0.1, None)]
        self.assertEqual(top_detection(dets).label, "青江菜")
        self.assertIsNone(top_detection([Detection("蔥", 0.1, None)]))
        self.assertIsNone(top_detection([]))

    def test_all_valid_vegetables_are_listed_once_by_confidence(self):
        from app.services.recognition import Detection, vegetable_detections
        dets = [Detection("高麗菜", 0.62, None), Detection("青江菜", 0.91, None), Detection("高麗菜", 0.8, None),
                Detection("", 0.99, None), Detection("番茄", float("nan"), None), Detection("蔥", 0.1, None)]
        found = vegetable_detections(dets)
        self.assertEqual([(d.label, d.confidence) for d in found], [("青江菜", 0.91), ("高麗菜", 0.8)])

    def test_label_map_covers_primary_model_and_uses_known_ids(self):
        import json
        import re
        from app.core.config import PROJECT_ROOT, settings
        from app.services.recognition import display_name
        data = json.loads(settings.recognition_labels_path.read_text(encoding="utf-8"))
        self.assertEqual(len(data["labels"]), 39)
        ids = [v["ingredient_id"] for v in data["labels"].values() if "ingredient_id" in v]
        self.assertTrue(all(re.fullmatch(r"FOOD\d{6}", i) for i in ids), ids)
        self.assertEqual(len(ids), len(set(ids)))
        workbook = PROJECT_ROOT / "data/recipe_dataset_1027.xlsx"
        if workbook.is_file():  # 原始工作簿不進 Git；本機有時核對編號確實存在於標準食材字典
            from openpyxl import load_workbook
            book = load_workbook(workbook, read_only=True, data_only=True)
            rows = book["標準食材整合"].iter_rows(values_only=True)
            next(rows)
            catalog = {row[1] for row in rows if row[1]}
            self.assertEqual([i for i in ids if i not in catalog], [])
        self.assertTrue(all(v["name"].strip() for v in data["labels"].values()))
        self.assertEqual(display_name("Cabbage"), ("高麗菜", "FOOD000867"))
        # 有歧義的類別不猜 ID；未列出的類別顯示原名
        self.assertEqual(display_name("garlic"), ("大蒜", None))
        self.assertEqual(display_name("unknown-class"), ("unknown-class", None))

    def test_malformed_service_response_is_unavailable(self):
        from app.services.recognition import RecognizerUnavailable, parse_output
        ok = parse_output({"model_version": "v1", "detections": [{"label": "高麗菜", "confidence": 0.8}]})
        self.assertEqual(ok.detections[0].label, "高麗菜")
        for bad in ({}, {"model_version": "", "detections": []},
                    {"model_version": "v1", "detections": [{"label": "x", "confidence": True}]},
                    {"model_version": "v1", "detections": "none"}):
            with self.assertRaises(RecognizerUnavailable):
                parse_output(bad)

    def test_http_adapter_distinguishes_failures(self):
        from app.services.recognition import HttpRecognizer, RecognizerUnavailable
        rec = HttpRecognizer("https://team-pc.invalid", TeamCredentials("token"), 1)
        req = httpx.Request("POST", "https://team-pc.invalid/v1/detect")
        cases = [
            (httpx.ConnectTimeout("t", request=req), "TIMEOUT"),
            (httpx.ConnectError("refused", request=req), "UNREACHABLE"),
            (httpx.Response(401, request=req), "AUTH_FAILED"),
            (httpx.Response(500, request=req), "SERVICE_ERROR"),
            (httpx.Response(200, request=req, content=b"not json"), "INVALID_RESPONSE"),
        ]
        for outcome, reason in cases:
            kwargs = {"side_effect": outcome} if isinstance(outcome, Exception) else {"return_value": outcome}
            with patch("httpx.post", **kwargs), self.assertRaises(RecognizerUnavailable) as ctx:
                rec.detect(b"jpeg")
            self.assertEqual(ctx.exception.reason, reason)
        good = httpx.Response(200, request=req, json={"model_version": "v8n-1", "detections": [
            {"label": "高麗菜", "confidence": 0.7, "box": [1, 2, 3, 4]}]})
        with patch("httpx.post", return_value=good) as post:
            out = rec.detect(b"jpeg")
        self.assertEqual(out.model_version, "v8n-1")
        self.assertEqual(post.call_args.kwargs["headers"], {"Authorization": "Bearer token"})


@requires_db
class T09ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        migrate_fresh()

    def setUp(self):
        truncate_all()
        seed_foods()
        from app.api.recognitions import get_recognizer
        from app.main import app
        self.app = app
        self.get_recognizer = get_recognizer
        self.recognizer = FakeRecognizer([("高麗菜", 0.62), ("青江菜", 0.91)])
        app.dependency_overrides[get_recognizer] = lambda: self.recognizer
        self.client = make_client()
        self.client.post("/auth/register", json={"email": "owner@example.test", "password": PASSWORD})
        self.client.post("/household", json={"name": "測試家庭"})

    def tearDown(self):
        self.app.dependency_overrides.pop(self.get_recognizer, None)

    def upload(self, data, filename="photo.jpg", client=None, **form):
        client = client or self.client
        return client.post("/recognitions", files={"image": (filename, data, "application/octet-stream")}, data=form)

    def test_every_detected_vegetable_is_returned(self):
        res = self.upload(jpeg_bytes())
        self.assertEqual(res.status_code, 201, res.text)
        body = res.json()
        self.assertEqual(body["status"], "succeeded")
        self.assertEqual([s["name"] for s in body["suggestions"]], ["青江菜", "高麗菜"])
        # 畫面不列信心或框線
        for hidden in ("detections", "confidence", "box", "candidates"):
            self.assertNotIn(hidden, res.text)
        from app.db.models import Recognition
        from db_support import session_factory
        with session_factory()() as db:
            row = db.get(Recognition, __import__("uuid").UUID(body["id"]))
        self.assertEqual(row.model_version, "yolov8n-synthetic-test")
        self.assertEqual(len(row.detections), 2)
        self.assertTrue(math.isclose(float(row.top_confidence), 0.91))

    def test_model_class_is_shown_in_chinese_with_standard_id(self):
        self.recognizer.detections = [("Cabbage", 0.8), ("garlic", 0.3)]
        body = self.upload(jpeg_bytes()).json()
        self.assertEqual(body["suggestions"][0], {"name": "高麗菜", "ingredient_id": "FOOD000867"})
        self.assertEqual(len(body["suggestions"]), 2)

    def test_no_valid_detection_asks_for_manual_input(self):
        self.recognizer.detections = [("蔥", 0.05)]
        body = self.upload(jpeg_bytes()).json()
        self.assertEqual((body["status"], body["suggestions"], body["retryable"]), ("no_result", [], False))
        self.assertIn("手動輸入", body["message"])

    def test_unavailable_service_keeps_photo_and_can_retry(self):
        from app.services.recognition import RecognizerUnavailable
        self.recognizer.error = RecognizerUnavailable("UNREACHABLE")
        body = self.upload(jpeg_bytes()).json()
        self.assertEqual((body["status"], body["retryable"]), ("unavailable", True))
        self.assertIn("照片已保留", body["message"])
        self.assertEqual(self.client.get(body["image_url"]).status_code, 200)
        self.recognizer.error = None
        retried = self.client.post(f"/recognitions/{body['id']}/retry").json()
        self.assertEqual((retried["status"], retried["suggestions"][0]["name"]), ("succeeded", "青江菜"))
        again = self.client.post(f"/recognitions/{body['id']}/retry")
        self.assertEqual(again.status_code, 409)

    def test_not_configured_service_is_unavailable_not_empty(self):
        self.app.dependency_overrides[self.get_recognizer] = lambda: None
        body = self.upload(jpeg_bytes()).json()
        self.assertEqual(body["status"], "unavailable")
        self.assertIn("手動輸入", body["message"])

    def test_image_is_decoded_rotated_and_stripped(self):
        body = self.upload(jpeg_bytes((64, 32), orientation=6, gps=True)).json()
        image = self.client.get(body["image_url"])
        self.assertEqual(image.headers["content-type"], "image/jpeg")
        stored = Image.open(io.BytesIO(image.content))
        self.assertEqual(stored.size, (32, 64))  # EXIF 方向 6 已轉正
        self.assertEqual(dict(stored.getexif()), {})  # 不保留 GPS 等中繼資料
        self.assertEqual(self.upload(png_bytes(), "a.png").status_code, 201)

    def test_invalid_uploads_are_rejected(self):
        from app.core.config import settings
        gif = io.BytesIO()
        Image.new("RGB", (5, 5)).save(gif, format="GIF")
        cases = [(b"not an image", "IMAGE_UNREADABLE"), (b"", "IMAGE_EMPTY"),
                 (gif.getvalue(), "IMAGE_FORMAT"), (jpeg_bytes()[:200], "IMAGE_UNREADABLE")]
        for data, reason in cases:
            res = self.upload(data)
            self.assertEqual(res.status_code, 422, reason)
            self.assertEqual(res.json()["error"]["details"]["reason"], reason)
        with patch.object(settings, "upload_max_bytes", 100):
            self.assertEqual(self.upload(jpeg_bytes()).json()["error"]["details"]["reason"], "IMAGE_TOO_LARGE")
        with patch.object(settings, "upload_max_pixels", 100):
            self.assertEqual(self.upload(jpeg_bytes()).json()["error"]["details"]["reason"], "IMAGE_TOO_MANY_PIXELS")
        self.assertEqual(self.recognizer.calls, 0)

    def test_confirm_records_corrected_type_without_weight(self):
        body = self.upload(jpeg_bytes()).json()
        bad = self.client.post(f"/recognitions/{body['id']}/confirm", json={"name": "高麗菜", "ingredient_id": "FOOD999999"})
        self.assertEqual(bad.status_code, 422)
        ok = self.client.post(f"/recognitions/{body['id']}/confirm",
                              json={"name": " 高麗菜 ", "ingredient_id": "FOOD000867"}).json()
        self.assertEqual(ok["confirmed"], [{"name": "高麗菜", "ingredient_id": "FOOD000867"}])
        weighted = self.client.post(f"/recognitions/{body['id']}/confirm",
                                    json={"name": "高麗菜", "ingredient_id": "FOOD000867", "weight_g": 300})
        self.assertEqual(weighted.status_code, 200)
        self.assertNotIn("weight", weighted.text)

    def test_one_photo_confirms_several_vegetables(self):
        body = self.upload(jpeg_bytes()).json()
        url = f"/recognitions/{body['id']}/confirm"
        self.client.post(url, json={"name": "青江菜", "ingredient_id": "FOOD000585"})
        self.client.post(url, json={"name": "甘藍", "ingredient_id": "FOOD000867"})
        again = self.client.post(url, json={"name": "高麗菜", "ingredient_id": "FOOD000867"}).json()
        self.assertEqual(again["confirmed"], [{"name": "青江菜", "ingredient_id": "FOOD000585"},
                                              {"name": "高麗菜", "ingredient_id": "FOOD000867"}])
        self.assertEqual(self.client.get(f"/recognitions/{body['id']}").json()["confirmed"], again["confirmed"])

    def test_delete_removes_photo(self):
        from app.services.recognition import image_file
        from app.db.models import Recognition
        from db_support import session_factory
        body = self.upload(jpeg_bytes()).json()
        with session_factory()() as db:
            path = image_file(db.get(Recognition, __import__("uuid").UUID(body["id"])).image_path)
        self.assertTrue(path.is_file())
        self.assertEqual(self.client.delete(f"/recognitions/{body['id']}").status_code, 204)
        self.assertFalse(path.exists())
        self.assertEqual(self.client.get(f"/recognitions/{body['id']}").status_code, 404)

    def test_cross_household_access_is_denied(self):
        body = self.upload(jpeg_bytes()).json()
        rid = body["id"]
        other = make_client()
        other.post("/auth/register", json={"email": "other@example.test", "password": PASSWORD})
        other.post("/household", json={"name": "別人家"})
        self.assertEqual(other.get(f"/recognitions/{rid}").status_code, 404)
        self.assertEqual(other.get(f"/recognitions/{rid}/image").status_code, 404)
        self.assertEqual(other.post(f"/recognitions/{rid}/retry").status_code, 404)
        self.assertEqual(other.post(f"/recognitions/{rid}/confirm",
                                    json={"name": "x", "ingredient_id": "FOOD000867"}).status_code, 404)
        self.assertEqual(other.delete(f"/recognitions/{rid}").status_code, 404)
        req = self.client.post("/meal-requests", json={}).json()
        foreign = self.upload(jpeg_bytes(), client=other, meal_request_id=req["id"])
        self.assertEqual(foreign.json()["error"]["details"]["reason"], "UNKNOWN_MEAL_REQUEST")
        anon = make_client()
        self.assertEqual(anon.get(f"/recognitions/{rid}/image").status_code, 401)


if __name__ == "__main__":
    unittest.main()
