"""同義食材對照：過敏全套用（含待確認）、家中已有與查價只用已確認組。"""
import json
import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch


def table(groups=(), broader=()):
    return {"groups": list(groups), "broader": list(broader)}


def group(gid, status, *ids):
    return {"id": gid, "status": status, "foods": [{"food_id": i, "name": i} for i in ids]}


class SynonymRuleTests(unittest.TestCase):
    def use(self, data):
        from app.core.config import settings
        from app.services import synonyms
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        path = Path(folder.name) / "synonyms.json"
        path.write_text(json.dumps(data), encoding="utf-8")
        patcher = patch.object(settings, "ingredient_synonyms_path", path)
        patcher.start()
        self.addCleanup(patcher.stop)
        synonyms._load.cache_clear()
        self.addCleanup(synonyms._load.cache_clear)
        return synonyms

    def test_pending_groups_expand_restrictions_both_ways(self):
        s = self.use(table([group("SYN1", "pending_confirmation", "TOMATO", "TOMATO_ALT")]))
        self.assertEqual([e["food_id"] for e in s.restriction_expansion(["TOMATO"])["TOMATO"]], ["TOMATO_ALT"])
        self.assertEqual([e["food_id"] for e in s.restriction_expansion(["TOMATO_ALT"])["TOMATO_ALT"]], ["TOMATO"])

    def test_rejected_groups_are_ignored(self):
        s = self.use(table([group("SYN1", "rejected", "A", "B")]))
        self.assertEqual(s.restriction_expansion(["A"]), {"A": []})

    def test_broader_category_applies_in_both_directions(self):
        broad = {"food_id": "BEEF", "status": "pending_confirmation",
                 "includes": [{"food_id": "STEAK"}, {"food_id": "BRISKET"}]}
        s = self.use(table(broader=[broad]))
        self.assertEqual({e["food_id"] for e in s.restriction_expansion(["BEEF"])["BEEF"]}, {"STEAK", "BRISKET"})
        # 對某個部位過敏：也排除只寫「牛肉」的食譜，但不擴大到其他部位
        self.assertEqual({e["food_id"] for e in s.restriction_expansion(["STEAK"])["STEAK"]}, {"BEEF"})

    def test_only_confirmed_groups_merge_available_ingredients(self):
        s = self.use(table([group("SYN1", "confirmed", "GARLIC", "GARLIC_ALT"),
                            group("SYN2", "pending_confirmation", "TOMATO", "TOMATO_ALT")]))
        self.assertEqual(s.confirmed_equivalents(["GARLIC"]), {"GARLIC", "GARLIC_ALT"})
        self.assertEqual(s.confirmed_equivalents(["TOMATO"]), {"TOMATO"})
        self.assertEqual(s.confirmed_groups(), [{"GARLIC", "GARLIC_ALT"}])

    def test_missing_table_means_no_synonyms(self):
        from app.core.config import settings
        from app.services import synonyms
        with patch.object(settings, "ingredient_synonyms_path", Path("does-not-exist.json")):
            synonyms._load.cache_clear()
            self.assertEqual(synonyms.restriction_expansion(["A"]), {"A": []})
        synonyms._load.cache_clear()


class ProjectTableTests(unittest.TestCase):
    def test_project_table_is_well_formed(self):
        from app.core.config import PROJECT_ROOT, settings
        data = json.loads(settings.ingredient_synonyms_path.read_text(encoding="utf-8"))
        ids = [f["food_id"] for g in data["groups"] for f in g["foods"]]
        self.assertEqual(len(ids), len(set(ids)), "同一食材不可出現在兩個同義組")
        self.assertTrue(all(re.fullmatch(r"FOOD\d{6}", i) for i in ids))
        self.assertTrue(all(len(g["foods"]) >= 2 and g["evidence"] for g in data["groups"]))
        self.assertTrue(all(g["status"] in {"pending_confirmation", "confirmed", "rejected"} for g in data["groups"]))
        workbook = PROJECT_ROOT / "data/recipe_dataset_1017.xlsx"
        if workbook.is_file():  # 本機有原始工作簿時核對編號與名稱
            from openpyxl import load_workbook
            rows = load_workbook(workbook, read_only=True, data_only=True)["標準食材整合"].iter_rows(values_only=True)
            next(rows)
            catalog = {row[1]: row[0] for row in rows if row[1]}
            for g in data["groups"]:
                for f in g["foods"]:
                    self.assertEqual(catalog.get(f["food_id"]), f["name"], f)


class PriceMergeTests(unittest.TestCase):
    def test_confirmed_synonyms_are_queried_once(self):
        from app.source_adapters.price_collection import merge_confirmed_synonyms
        from app.source_adapters.price_models import IngredientQuery
        queries = [IngredientQuery(canonical_name="大蒜", food_ids=["G1"], aliases=["蒜"]),
                   IngredientQuery(canonical_name="番茄", food_ids=["T1"]),
                   IngredientQuery(canonical_name="蒜頭", food_ids=["G2"])]
        merged = merge_confirmed_synonyms(queries, [{"G1", "G2"}])
        self.assertEqual([q.canonical_name for q in merged], ["大蒜", "番茄"])
        self.assertEqual(merged[0].food_ids, ["G1", "G2"])
        self.assertIn("蒜頭", merged[0].aliases)
        self.assertEqual(merge_confirmed_synonyms(queries, []), queries)


if __name__ == "__main__":
    unittest.main()
