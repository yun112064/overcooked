"""T01：食譜匯入的品質規則與重跑行為（AC05、G01）。工作簿皆為合成資料。"""
import tempfile
import unittest
from pathlib import Path

from sqlalchemy import func, select

from db_support import migrate_fresh, requires_db, session_factory, truncate_all
from recipe_workbook import build, ingredient, recipe


SOURCE_RIGHTS = {"合成來源": {"license_name": "合成來源授權（測試）", "rights_status": "verified",
                            "source_version": "v1", "retrieved_at": "2026-10-17T00:00:00+08:00"}}


def parse(recipes, catalog=None, source_rights=None, review_record=None):
    from app.services.recipe_import import parse_recipe_workbook
    with tempfile.TemporaryDirectory() as folder:
        return parse_recipe_workbook(build(Path(folder) / "book.xlsx", recipes, catalog), source_rights, review_record)


def quantity(raw_text, amount, unit="g", name="高麗菜"):
    bundle = parse([recipe("R001", "清炒高麗菜", ingredients=[
        ingredient("FOOD000005", name, amount, raw_text=raw_text, unit=unit, weight_g=amount)])])
    return bundle.recipe_ingredients[0], by_id(bundle)["R001"]


def by_id(bundle):
    return {r["recipe_id"]: r for r in bundle.recipes}


class QualityRuleTests(unittest.TestCase):
    def test_fully_reviewed_recipe_qualifies(self):
        bundle = parse([recipe("R001", "清炒高麗菜")])
        self.assertTrue(by_id(bundle)["R001"]["is_eligible"], by_id(bundle)["R001"]["quality_issues"])
        self.assertEqual(bundle.report["g01"]["actual_qualified_recipes"], 1)

    def test_nonpositive_servings_is_stored_null_and_blocks(self):
        row = by_id(parse([recipe("R001", "清炒高麗菜", servings="0人份")]))["R001"]
        self.assertIsNone(row["original_servings"])
        self.assertEqual(row["servings_raw"], "0人份")
        self.assertIn("nonpositive_servings", row["quality_issues"])
        self.assertFalse(row["is_eligible"])

    def test_inverted_adjustable_range_is_stored_null_and_blocks(self):
        bundle = parse([recipe("R001", "清炒高麗菜", ingredients=[
            ingredient("FOOD000005", "高麗菜", 300, adjustable_min=400, adjustable_max=200)])])
        row = bundle.recipe_ingredients[0]
        self.assertEqual((row["adjustable_min"], row["adjustable_max"]), (None, None))
        self.assertFalse(by_id(bundle)["R001"]["is_eligible"])
        self.assertIn("ingredient:invalid_adjustable_range", by_id(bundle)["R001"]["quality_issues"])

    def test_bad_numeric_cell_only_affects_that_recipe(self):
        bundle = parse([
            recipe("R001", "清炒高麗菜", ingredients=[ingredient("FOOD000005", "高麗菜", 300, amount="適量", weight_g="很多")]),
            recipe("R002", "蒜炒青江菜", ingredients=[ingredient("FOOD000006", "青江菜", 250)]),
        ])
        rows = by_id(bundle)
        bad = bundle.recipe_ingredients[0]
        self.assertEqual((bad["source_amount"], bad["source_weight_g"]), (None, None))
        self.assertFalse(rows["R001"]["is_eligible"])
        self.assertIn("ingredient:invalid_numeric_value:amount", rows["R001"]["quality_issues"])
        self.assertTrue(rows["R002"]["is_eligible"], rows["R002"]["quality_issues"])

    def test_negative_amount_is_rejected_per_row(self):
        bundle = parse([recipe("R001", "清炒高麗菜", ingredients=[ingredient("FOOD000005", "高麗菜", -5)])])
        self.assertIsNone(bundle.recipe_ingredients[0]["amount"])
        self.assertFalse(by_id(bundle)["R001"]["is_eligible"])

    def test_exact_duplicates_count_once(self):
        bundle = parse([recipe("R002", "番茄炒蛋"), recipe("R001", "番茄炒蛋"), recipe("R003", "清炒高麗菜")])
        rows = by_id(bundle)
        self.assertTrue(rows["R001"]["is_eligible"])
        self.assertFalse(rows["R002"]["is_eligible"])
        self.assertIn("duplicate_not_primary", rows["R002"]["quality_issues"])
        self.assertEqual(rows["R002"]["duplicate_of_recipe_id"], "R001")
        self.assertEqual(bundle.report["g01"]["actual_qualified_recipes"], 2)

    def test_import_selection_keeps_primary_and_reports_excel_rows(self):
        from app.services.recipe_import import select_unique_recipe_bundle
        original = parse([recipe("R002", "番茄炒蛋"), recipe("R001", "番茄炒蛋")])
        selected = select_unique_recipe_bundle(original)
        self.assertEqual([row["recipe_id"] for row in selected.recipes], ["R001"])
        self.assertEqual([row["raw_id"] for row in selected.source_archive], ["RAW-R001"])
        self.assertTrue(all(row["recipe_id"] == "R001" for row in selected.recipe_steps))
        self.assertTrue(all(row["recipe_id"] == "R001" for row in selected.recipe_ingredients))
        choice = selected.report["import_selection"]["selections"][0]
        self.assertEqual((choice["kept"]["recipes_row"], choice["kept"]["source_archive_row"]), (3, 3))
        self.assertEqual(choice["skipped"][0]["recipe_id"], "R002")
        self.assertEqual(len(original.recipes), 2)  # Raw validation evidence is retained.
        self.assertEqual(select_unique_recipe_bundle(selected), selected)

    def test_import_selection_does_not_remove_suspected_duplicates(self):
        from app.services.recipe_import import select_unique_recipe_bundle
        same = [ingredient("FOOD000001", "番茄", 200), ingredient("FOOD000002", "雞蛋", 150)]
        selected = select_unique_recipe_bundle(parse([
            recipe("R001", "番茄炒蛋", ingredients=same),
            recipe("R002", "家常番茄炒蛋", ingredients=same)]))
        self.assertEqual(len(selected.recipes), 2)
        self.assertEqual(selected.report["import_selection"]["skipped_duplicate_recipes"], 0)

    def test_suspected_duplicates_count_once_until_reviewed(self):
        same = [ingredient("FOOD000001", "番茄", 200), ingredient("FOOD000002", "雞蛋", 150)]
        bundle = parse([recipe("R001", "番茄炒蛋", ingredients=same),
                        recipe("R002", "家常番茄炒蛋", ingredients=same)])
        rows = by_id(bundle)
        self.assertEqual(rows["R001"]["duplicate_status"], "suspected_duplicate")
        self.assertTrue(rows["R001"]["is_eligible"])
        self.assertFalse(rows["R002"]["is_eligible"])
        self.assertIn("suspected_duplicate_not_primary", rows["R002"]["quality_issues"])

    def test_catalog_sheet_becomes_the_food_dictionary(self):
        bundle = parse([recipe("R001", "清炒高麗菜")],
                       catalog={"FOOD000005": "高麗菜", "FOOD001566": "雞蛋", "FOOD000039": "大蒜"})
        aliases = {row["food_id"]: row["canonical_name"] for row in bundle.ingredient_aliases}
        self.assertEqual(aliases, {"FOOD000005": "高麗菜", "FOOD001566": "雞蛋", "FOOD000039": "大蒜"})
        self.assertEqual(bundle.report["counts"]["catalog_food_ids"], 3)

    def test_recipe_food_missing_from_catalog_is_rejected(self):
        from app.services.recipe_import import WorkbookValidationError
        with self.assertRaises(WorkbookValidationError):
            parse([recipe("R001", "清炒高麗菜")], catalog={"FOOD001566": "雞蛋"})
        with self.assertRaises(WorkbookValidationError):  # 同一編號名稱不一致
            parse([recipe("R001", "清炒高麗菜")], catalog={"FOOD000005": "甘藍菜"})

    def test_source_rights_config_fills_missing_rights_only(self):
        no_rights = [recipe("R001", "清炒高麗菜", rights_in_workbook=False)]
        before = by_id(parse(no_rights))["R001"]
        self.assertIn("missing_license", before["quality_issues"])
        self.assertFalse(before["is_eligible"])
        after = by_id(parse(no_rights, source_rights=SOURCE_RIGHTS))["R001"]
        for issue in ("missing_license", "rights_not_verified", "missing_source_version", "missing_retrieved_at"):
            self.assertNotIn(issue, after["quality_issues"])
        self.assertTrue(after["is_eligible"], after["quality_issues"])
        # 授權不能取代人工審核
        pending = by_id(parse([recipe("R001", "清炒高麗菜", rights_in_workbook=False, approved=False)],
                              source_rights=SOURCE_RIGHTS))["R001"]
        self.assertIn("recipe_review_pending", pending["quality_issues"])
        self.assertFalse(pending["is_eligible"])

    def test_vague_words_use_reference_amount(self):
        for raw in ("鹽 適量", "胡椒 少許", "雞腿 一大隻", "香菇 十幾朵", "鱸魚 數塊", "醬油"):
            row, recipe_row = quantity(raw, 10, name=raw.split()[0])
            self.assertEqual((row["quantity_status"], row["amount"]), ("reference_amount", 10), raw)
            self.assertTrue(recipe_row["is_eligible"], (raw, recipe_row["quality_issues"]))

    def test_reference_amount_needs_an_amount(self):
        row, recipe_row = quantity("鹽 適量", None, name="鹽")
        self.assertEqual(row["quantity_status"], "qualitative_unresolved")
        self.assertFalse(recipe_row["is_eligible"])

    def test_chinese_numerals_must_match_amount(self):
        cases = {"雞蛋 六顆": 6, "醬油露 半碗": 0.5, "高麗菜 一顆半": 1.5, "肉 十二片": 12, "米 兩杯": 2}
        for raw, value in cases.items():
            row, _ = quantity(raw, value, unit="顆", name=raw.split()[0])
            self.assertEqual(row["quantity_status"], "source_exact", raw)
        # 原文與 amount 不一致時以 Excel 的 amount 為準，只留下提醒、不阻擋
        row, recipe_row = quantity("雞蛋 六顆", 2, unit="顆", name="雞蛋")
        self.assertEqual((row["quantity_status"], row["amount"]), ("reference_amount", 2))
        self.assertIn("ingredient:raw_text_differs_from_amount", recipe_row["quality_issues"])
        self.assertTrue(recipe_row["is_eligible"], recipe_row["quality_issues"])

    def test_notes_and_names_are_not_read_as_quantities(self):
        row, _ = quantity("[內餡食材(六份)] 鮪魚罐頭 一罐", 1, unit="罐", name="鮪魚罐頭")
        self.assertEqual(row["quantity_status"], "source_exact")
        row, _ = quantity("九層塔(摘葉子就好) 半包", 0.5, unit="包", name="九層塔")
        self.assertEqual(row["quantity_status"], "source_exact")

    def test_review_record_applies_only_to_the_reviewed_workbook(self):
        from app.services.recipe_import import parse_recipe_workbook
        books = [recipe("R001", "清炒高麗菜", approved=None, rights_in_workbook=False)]
        with tempfile.TemporaryDirectory() as folder:
            path = build(Path(folder) / "book.xlsx", books)
            import hashlib
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            approvals = {key: "approved" for key in
                         ("source_review", "recipe_review", "cooking_method", "step_review", "adjustable_range")}
            record = {"workbook_sha256": digest, "reviewed_on": "2026-10-19", "approvals": approvals}
            ok = parse_recipe_workbook(path, SOURCE_RIGHTS, record)
            stale = parse_recipe_workbook(path, SOURCE_RIGHTS, {**record, "workbook_sha256": "0" * 64})
        self.assertTrue(by_id(ok)["R001"]["is_eligible"], by_id(ok)["R001"]["quality_issues"])
        self.assertTrue(ok.report["review_record"]["applied"])
        self.assertFalse(by_id(stale)["R001"]["is_eligible"])
        self.assertFalse(stale.report["review_record"]["applied"])

    def test_reviewed_distinct_pair_is_not_a_duplicate_for_that_workbook_only(self):
        from app.services.recipe_import import parse_recipe_workbook
        same = [ingredient("FOOD000001", "番茄", 200), ingredient("FOOD000002", "雞蛋", 150)]
        books = [recipe(recipe_id, name, ingredients=same, approved=None, rights_in_workbook=False)
                 for recipe_id, name in (("R001", "番茄炒蛋"), ("R002", "家常番茄炒蛋"))]
        with tempfile.TemporaryDirectory() as folder:
            path = build(Path(folder) / "book.xlsx", books)
            import hashlib
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            approvals = {key: "approved" for key in
                         ("source_review", "recipe_review", "cooking_method", "step_review", "adjustable_range")}
            record = {"workbook_sha256": digest, "approvals": approvals, "distinct_recipe_pairs": [["R002", "R001"]]}
            reviewed = parse_recipe_workbook(path, SOURCE_RIGHTS, record)
            stale = parse_recipe_workbook(path, SOURCE_RIGHTS, {**record, "workbook_sha256": "0" * 64})
        rows = by_id(reviewed)
        self.assertTrue(rows["R002"]["is_eligible"], rows["R002"]["quality_issues"])
        self.assertEqual(rows["R002"]["duplicate_status"], "unique_candidate")
        self.assertEqual(reviewed.report["duplicate_review"]["suspected_pairs"][0]["status"], "reviewed_distinct")
        self.assertIn("suspected_duplicate_review_required", by_id(stale)["R002"]["quality_issues"])

    def test_reviewed_cooking_method_fills_recipe_without_method_keywords(self):
        from app.services.recipe_import import WorkbookValidationError, parse_recipe_workbook
        books = [recipe("R001", "涼拌小黃瓜", steps=("小黃瓜拍開切段，加醋靜置",), approved=None,
                        rights_in_workbook=False)]
        with tempfile.TemporaryDirectory() as folder:
            path = build(Path(folder) / "book.xlsx", books)
            from openpyxl import load_workbook
            book = load_workbook(path)  # 如真實工作簿：recipes 未填 cooking_methods
            sheet = book["recipes"]
            column = [cell.value for cell in sheet[1]].index("cooking_methods") + 1
            sheet.cell(row=2, column=column).value = None
            book.save(path)
            import hashlib
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            approvals = {key: "approved" for key in
                         ("source_review", "recipe_review", "cooking_method", "step_review", "adjustable_range")}
            record = {"workbook_sha256": digest, "approvals": approvals, "recipe_cooking_methods": {"R001": ["拌"]}}
            reviewed = parse_recipe_workbook(path, SOURCE_RIGHTS, record)
            without = parse_recipe_workbook(path, SOURCE_RIGHTS, {**record, "recipe_cooking_methods": {}})
            with self.assertRaises(WorkbookValidationError):
                parse_recipe_workbook(path, SOURCE_RIGHTS, {**record, "recipe_cooking_methods": {"R001": ["醃"]}})
        self.assertEqual(by_id(reviewed)["R001"]["cooking_methods"], ["拌"])
        self.assertTrue(by_id(reviewed)["R001"]["is_eligible"], by_id(reviewed)["R001"]["quality_issues"])
        self.assertIn("missing_cooking_method", by_id(without)["R001"]["quality_issues"])

    def test_unreviewed_recipe_never_pads_the_count(self):
        bundle = parse([recipe("R001", "清炒高麗菜", approved=False)])
        self.assertFalse(by_id(bundle)["R001"]["is_eligible"])
        self.assertEqual(bundle.report["g01"]["status"], "not_passed")
        self.assertEqual(bundle.report["g01"]["gap"], 150)


@requires_db
class ImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        migrate_fresh()

    def setUp(self):
        truncate_all()
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)

    def run_import(self, recipes, name="book.xlsx"):
        from app.services.recipe_import import import_recipe_bundle, parse_recipe_workbook
        bundle = parse_recipe_workbook(build(Path(self.folder.name) / name, recipes))
        with session_factory()() as session, session.begin():
            return import_recipe_bundle(session, bundle)

    def count(self, model, **where):
        with session_factory()() as session:
            query = select(func.count()).select_from(model)
            for key, value in where.items():
                query = query.where(getattr(model, key) == value)
            return session.scalar(query)

    def test_invalid_values_no_longer_roll_back_the_whole_batch(self):
        from app.models import Recipe, RecipeIngredient
        self.run_import([
            recipe("R001", "清炒高麗菜", servings="0人份"),
            recipe("R002", "蒜炒青江菜", ingredients=[
                ingredient("FOOD000006", "青江菜", 250, adjustable_min=300, adjustable_max=100)]),
            recipe("R003", "番茄炒蛋"),
        ])
        self.assertEqual(self.count(Recipe), 3)
        self.assertEqual(self.count(Recipe, is_eligible=True), 1)
        self.assertEqual(self.count(RecipeIngredient), 3)

    def test_duplicate_listed_before_primary_only_imports_primary(self):
        from app.models import Recipe, SourceArchive, RecipeIngredient, RecipeStep
        self.run_import([recipe("R002", "番茄炒蛋"), recipe("R001", "番茄炒蛋")])
        with session_factory()() as session:
            self.assertIsNone(session.get(Recipe, "R002"))
            self.assertIsNotNone(session.get(Recipe, "R001"))
        for model in (Recipe, SourceArchive, RecipeIngredient, RecipeStep):
            self.assertEqual(self.count(model), 1)

    def test_reimport_is_idempotent(self):
        from app.models import RecipeIngredient, RecipeStep
        books = [recipe("R001", "清炒高麗菜", steps=("切菜", "下鍋炒"))]
        first = self.run_import(books, "a.xlsx")
        second = self.run_import(books, "b.xlsx")
        self.assertEqual(first["recipe_steps"], {"inserted": 2, "updated": 0, "deleted": 0})
        self.assertEqual(second["recipe_steps"], {"inserted": 0, "updated": 2, "deleted": 0})
        self.assertEqual((self.count(RecipeStep), self.count(RecipeIngredient)), (2, 1))

    def test_reimport_removes_rows_dropped_from_workbook(self):
        from app.models import RecipeIngredient, RecipeStep
        self.run_import([recipe("R001", "番茄炒蛋", steps=("切番茄", "炒蛋", "拌炒"), ingredients=[
            ingredient("FOOD000001", "番茄", 200), ingredient("FOOD000002", "雞蛋", 150)])], "a.xlsx")
        result = self.run_import([recipe("R001", "番茄炒蛋", steps=("切番茄", "拌炒"), ingredients=[
            ingredient("FOOD000001", "番茄", 200)])], "b.xlsx")
        self.assertEqual(result["recipe_steps"]["deleted"], 1)
        self.assertEqual(result["recipe_ingredients"]["deleted"], 1)
        self.assertEqual((self.count(RecipeStep), self.count(RecipeIngredient)), (2, 1))

    def test_reimport_deletes_recipes_removed_from_workbook(self):
        from app.models import Recipe, RecipeIngredient, RecipeStep, SourceArchive
        self.run_import([recipe("R001", "清炒高麗菜"), recipe("R002", "番茄炒蛋", steps=("切", "炒"), ingredients=[
            ingredient("FOOD000001", "番茄", 200)])], "a.xlsx")
        result = self.run_import([recipe("R001", "清炒高麗菜")], "b.xlsx")
        self.assertEqual(result["recipes"]["deleted"], 1)
        self.assertEqual(result["source_archive"]["deleted"], 1)
        self.assertEqual(result["recipe_steps"]["deleted"], 2)
        self.assertEqual(result["recipe_ingredients"]["deleted"], 1)
        self.assertEqual(self.count(Recipe), 1)
        self.assertEqual(self.count(RecipeStep, recipe_id="R002"), 0)
        self.assertEqual(self.count(RecipeIngredient, recipe_id="R002"), 0)
        self.assertEqual(self.count(SourceArchive), 1)
        # 刪除後再加回來也能正常匯入
        again = self.run_import([recipe("R001", "清炒高麗菜"), recipe("R002", "番茄炒蛋")], "c.xlsx")
        self.assertEqual((again["recipes"]["inserted"], again["recipes"]["deleted"]), (1, 0))

    def test_reimport_with_renumbered_steps(self):
        from app.models import RecipeStep
        base = recipe("R001", "番茄炒蛋", steps=("切番茄", "炒蛋"))
        self.run_import([base], "a.xlsx")
        swapped = {**base, "step_ids": {1: "R001-S2", 2: "R001-S1"}}  # 兩個步驟交換編號
        self.run_import([swapped], "b.xlsx")
        with session_factory()() as session:
            numbers = dict(session.execute(select(RecipeStep.step_id, RecipeStep.step_no)).all())
        self.assertEqual(numbers, {"R001-S1": 2, "R001-S2": 1})


if __name__ == "__main__":
    unittest.main()
