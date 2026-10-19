"""產生 T01 六工作表的合成食譜工作簿（測試用，非真實食譜來源）。

recipe(...) 預設產生「所有審核都通過」的食譜，方便測試只改一個條件。
"""
from pathlib import Path

from openpyxl import Workbook

SOURCE_COLUMNS = ["raw_id", "source_name", "source_url", "source_recipe_id", "recipe_name_raw", "servings_raw",
                  "ingredients_raw", "steps_raw", "tags_raw", "license_name", "rights_status", "source_version",
                  "retrieved_at", "review_status"]
RECIPE_COLUMNS = ["raw_id", "recipe_id", "recipe_name", "servings_raw", "category", "meal_type", "Cuisine",
                  "source_name", "source_url", "allergens", "cooking_methods", "cooking_method_status", "review_status"]
STEP_COLUMNS = ["step_id", "recipe_id", "step_no", "instruction_clean"]
INGREDIENT_COLUMNS = ["recipe_ingredient_id", "recipe_id", "sequence", "raw_text", "canonical_name", "Food_ID",
                      "amount", "unit", "weight_g", "protein_g", "carbohydrates_g", "fat_g", "sugar_g", "sodium_mg",
                      "allergens", "adjustable_min", "adjustable_max", "adjustment_status"]


def ingredient(food_id, name, grams, **extra):
    row = {"food_id": food_id, "name": name, "raw_text": f"{name} {grams}g", "amount": grams, "unit": "g",
           "weight_g": grams, "adjustable_min": grams, "adjustable_max": grams, "adjustment_status": "approved"}
    row.update(extra)
    return row


def recipe(recipe_id, name, *, servings="2人份", steps=("熱鍋下油，大火快炒",), ingredients=None, approved=True,
           rights_in_workbook=True):
    """rights_in_workbook=False：來源列不附授權欄位（如真實工作簿），改由來源設定補上。"""
    return {"id": recipe_id, "name": name, "servings": servings, "steps": list(steps), "approved": approved,
            "rights_in_workbook": rights_in_workbook,
            "ingredients": ingredients or [ingredient("FOOD000005", "高麗菜", 300)]}


def build(path: Path, recipes: list[dict], catalog: dict[str, str] | None = None) -> Path:
    """catalog：選用的「標準食材整合」工作表內容 {Food_ID: 標準名稱}。"""
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {name: wb.create_sheet(name) for name in
              ("source_archive", "recipes", "recipe_steps", "recipe_ingredients", "ingredient_aliases", "allergens")}
    sheets["source_archive"].append(SOURCE_COLUMNS)
    sheets["recipes"].append(RECIPE_COLUMNS)
    sheets["recipe_steps"].append(STEP_COLUMNS)
    sheets["recipe_ingredients"].append(INGREDIENT_COLUMNS)
    sheets["ingredient_aliases"].append(["alias_id", "Food_ID", "canonical_name", "alias"])
    sheets["allergens"].append(["allergen_id", "allergen_name", "reference"])
    foods = {}
    for r in recipes:
        ok = "approved" if r["approved"] else "pending"
        raw_id = f"RAW-{r['id']}"
        rights = (["CC BY 4.0（合成測試）", "verified" if r["approved"] else "unknown", "v1", "2026-10-01T00:00:00Z"]
                  if r["rights_in_workbook"] else [None, None, None, None])
        sheets["source_archive"].append([
            raw_id, "合成來源", f"https://example.invalid/{r['id']}", r["id"], r["name"], r["servings"],
            "、".join(i["name"] for i in r["ingredients"]), "；".join(r["steps"]), "家常", *rights, ok])
        sheets["recipes"].append([
            raw_id, r["id"], r["name"], r["servings"], "主菜", "晚餐", "台式", "合成來源",
            f"https://example.invalid/{r['id']}", "", "炒", ok, ok])
        for no, text in enumerate(r["steps"], start=1):
            step_id = r.get("step_ids", {}).get(no, f"{r['id']}-S{no}")
            sheets["recipe_steps"].append([step_id, r["id"], no, text])
        for seq, i in enumerate(r["ingredients"], start=1):
            foods[i["food_id"]] = i["name"]
            sheets["recipe_ingredients"].append([
                f"{r['id']}-I{seq}", r["id"], seq, i["raw_text"], i["name"], i["food_id"], i["amount"], i["unit"],
                i["weight_g"], 1, 1, 1, 1, 1, "", i["adjustable_min"], i["adjustable_max"], i["adjustment_status"]])
    for n, (food_id, name) in enumerate(sorted(foods.items()), start=1):
        sheets["ingredient_aliases"].append([f"A{n}", food_id, name, name])
    sheets["allergens"].append(["AL1", "蛋", "合成參考"])
    if catalog is not None:
        sheet = wb.create_sheet("標準食材整合")
        sheet.append(["合併標準食材名稱", "Food_ID", None])
        for food_id, name in catalog.items():
            sheet.append([name, food_id, None])
        sheet.append([None, None, None])  # 真實工作簿尾端有空白列
    wb.save(path)
    return path
