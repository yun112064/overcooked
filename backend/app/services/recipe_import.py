from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, replace
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.models.recipe import (
    Allergen,
    IngredientAlias,
    Recipe,
    RecipeIngredient,
    RecipeStep,
    SourceArchive,
)


SCHEMA_VERSION = "t01-recipe-import-v1"

REQUIRED_COLUMNS = {
    "source_archive": {
        "raw_id",
        "source_name",
        "source_url",
        "source_recipe_id",
        "recipe_name_raw",
        "servings_raw",
        "ingredients_raw",
        "steps_raw",
        "tags_raw",
    },
    "recipes": {
        "raw_id",
        "recipe_id",
        "recipe_name",
        "servings_raw",
        "category",
        "meal_type",
        "Cuisine",
        "source_name",
        "source_url",
        "allergens",
    },
    "recipe_steps": {"step_id", "recipe_id", "step_no", "instruction_clean"},
    "recipe_ingredients": {
        "recipe_ingredient_id",
        "recipe_id",
        "sequence",
        "raw_text",
        "canonical_name",
        "Food_ID",
        "amount",
        "unit",
        "weight_g",
        "protein_g",
        "carbohydrates_g",
        "fat_g",
        "sugar_g",
        "sodium_mg",
        "allergens",
    },
    "ingredient_aliases": {"alias_id", "Food_ID", "canonical_name", "alias"},
    "allergens": {"allergen_id", "allergen_name", "reference"},
}

# 選用工作表：全系統共用的標準食材字典（Food_ID ↔ 標準名稱）
CATALOG_SHEET = "標準食材整合"
CATALOG_COLUMNS = {"合併標準食材名稱", "Food_ID"}

QUALITATIVE_WORDS = (
    "適量",
    "少許",
    "酌量",
    "若干",
    "隨意",
    "一點",
    "些許",
    "適當",
    "少量",
    "視情況",
    "依喜好",
    "任意",
)

METHOD_KEYWORDS = (
    "氣炸",
    "微波",
    "煎",
    "炒",
    "炸",
    "烤",
    "蒸",
    "煮",
    "燉",
    "滷",
    "燙",
    "拌",
)

UNIT_ALIASES = {
    "公克": "g",
    "克": "g",
    "G": "g",
    "毫升": "ml",
    "cc": "ml",
    "CC": "ml",
    "茶匙": "小匙",
    "湯匙": "大匙",
    "枝": "支",
}

VALID_UNITS = {
    "g",
    "ml",
    "L",
    "公升",
    "大匙",
    "小匙",
    "匙",
    "顆",
    "片",
    "條",
    "根",
    "瓣",
    "支",
    "個",
    "碗",
    "隻",
    "杯",
    "把",
    "朵",
    "盒",
    "包",
    "斤",
    "塊",
    "粒",
    "小塊",
    "株",
    "罐",
    "尾",
    "人份",
    "瓶",
    "份",
    "兩",
    "張",
    "米杯",
    "小撮",
    "滴",
    "小把",
    "小飯碗",
    "大把",
    "段",
    "蓋",
    "小粒",
    "對",
    "大株",
    "大瓣",
    "大片",
}

RANGE_RE = re.compile(r"(?P<low>\d+(?:\.\d+)?)\s*[~～\-–—]\s*(?P<high>\d+(?:\.\d+)?)")
FRACTION_RE = re.compile(r"(?P<num>\d+)\s*/\s*(?P<den>\d+)")
NUMBER_RE = re.compile(r"\d+(?:\.\d+)?")
# 原文以中文數字寫的數量（例如「雞蛋 六顆」「醬油露 半碗」）
CHINESE_NUMBER_RE = re.compile(r"(?P<num>[一二兩三四五六七八九十]+)(?P<half>半)?|(?P<only_half>半)")
# 原文數量含這些字時屬於概略用量（「一大隻」「十幾朵」「數塊」「六到七瓣」）
VAGUE_QUANTITY_RE = re.compile(r"[大小幾數約多到至]|左右")
NOTE_RE = re.compile(r"\[[^\]]*\]|【[^】]*】|\([^)]*\)|（[^）]*）")
CHINESE_DIGITS = {"一": 1, "二": 2, "兩": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
UNIT_IN_TEXT_RE = re.compile(
    r"(?:\d+(?:\.\d+)?\s*(?:/\s*\d+)?|\d+(?:\.\d+)?\s*[~～\-–—]\s*\d+(?:\.\d+)?)\s*"
    r"(?P<unit>公克|克|g|G|毫升|ml|cc|CC|大匙|湯匙|小匙|茶匙|顆|片|條|根|瓣|支|個|碗|隻|杯|把|朵|盒|包|斤|塊|粒|株|罐|尾|份|兩|張|米杯|小撮|滴)"
)


class WorkbookValidationError(ValueError):
    """Raised when the workbook contract is incomplete or inconsistent."""


@dataclass(frozen=True)
class ImportBundle:
    source_archive: list[dict[str, Any]]
    recipes: list[dict[str, Any]]
    recipe_steps: list[dict[str, Any]]
    recipe_ingredients: list[dict[str, Any]]
    ingredient_aliases: list[dict[str, Any]]
    allergens: list[dict[str, Any]]
    report: dict[str, Any]


def _text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def _decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise WorkbookValidationError(f"Invalid numeric value: {value!r}") from exc


def _lenient_decimal(value: Any, field: str, issues: list[str]) -> Decimal | None:
    """食材數值欄位：非數字或負數存成 NULL 並記錄原因，不中斷整批匯入。

    結構欄位（序號、步驟編號等）仍用 _decimal／_integer 嚴格檢查。
    """
    try:
        number = _decimal(value)
    except WorkbookValidationError:
        issues.append(f"invalid_numeric_value:{field}")
        return None
    if number is not None and (not number.is_finite() or number < 0):
        issues.append(f"invalid_numeric_value:{field}")
        return None
    return number


def _integer(value: Any, field: str) -> int:
    number = _decimal(value)
    if number is None or number != number.to_integral_value():
        raise WorkbookValidationError(f"{field} must be an integer: {value!r}")
    return int(number)


def _split_list(value: Any) -> list[str]:
    text_value = _text(value)
    if text_value is None:
        return []
    return [item.strip() for item in re.split(r"[；;]", text_value) if item.strip()]


def normalize_name(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value).lower()
    return "".join(character for character in normalized if character.isalnum())


def standard_dish_id(recipe_name: str) -> str:
    normalized = normalize_name(recipe_name)
    digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:12].upper()
    return f"DISH{digest}"


def _content_hash(record: dict[str, Any]) -> str:
    payload = json.dumps(record, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sheet_records(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise WorkbookValidationError(f"Missing worksheet: {sheet_name}")
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise WorkbookValidationError(f"Worksheet is empty: {sheet_name}") from exc

    headers = [_text(value) for value in header_row]
    if any(header is None for header in headers):
        raise WorkbookValidationError(f"Worksheet has blank header: {sheet_name}")
    missing = REQUIRED_COLUMNS[sheet_name] - set(headers)
    if missing:
        raise WorkbookValidationError(
            f"Worksheet {sheet_name} missing columns: {', '.join(sorted(missing))}"
        )

    records: list[dict[str, Any]] = []
    for row_number, values in enumerate(rows, start=2):
        if all(value is None or value == "" for value in values):
            continue
        record = {str(header): value for header, value in zip(headers, values)}
        record["__row__"] = row_number
        records.append(record)
    return records


def _parse_servings(value: Any) -> tuple[Decimal | None, str, list[str]]:
    raw = _text(value)
    if raw is None:
        return None, "missing", ["missing_servings"]
    match = NUMBER_RE.search(raw)
    if match is None:
        return None, "unresolved", ["unresolved_servings"]
    servings = Decimal(match.group())
    issues = []
    status = "parsed"
    if servings <= 0:
        # 0 或負數違反資料表 CHECK；存成 NULL 並列為阻擋原因，避免整批匯入回滾
        issues.append("nonpositive_servings")
        return None, "invalid", issues
    elif servings > 12:
        issues.append("servings_outlier_review_required")
        status = "review_required"
    return servings, status, issues


def _normalize_unit(source_unit: str | None, raw_text: str) -> tuple[str | None, str, list[str]]:
    issues: list[str] = []
    if source_unit is None:
        return None, "missing", ["missing_unit"]
    candidate = UNIT_ALIASES.get(source_unit, source_unit)
    if candidate in VALID_UNITS:
        status = "normalized" if candidate != source_unit else "source"
        return candidate, status, issues

    match = UNIT_IN_TEXT_RE.search(raw_text)
    if match:
        parsed = UNIT_ALIASES.get(match.group("unit"), match.group("unit"))
        if parsed in VALID_UNITS:
            issues.append("unit_corrected_from_raw_text")
            return parsed, "corrected_from_raw", issues
    issues.append("invalid_unit")
    return None, "invalid", issues


def _chinese_number(text: str) -> Decimal | None:
    """解析原文中的中文數字（一～九十九、兩、半、X半）；無法解析時回傳 None。"""
    match = CHINESE_NUMBER_RE.search(text)
    if match is None:
        return None
    if match.group("only_half"):
        return Decimal("0.5")
    digits = match.group("num")
    if "十" in digits:
        tens, _, ones = digits.partition("十")
        if (tens and tens not in CHINESE_DIGITS) or (ones and ones not in CHINESE_DIGITS) or "十" in ones:
            return None
        value = Decimal(CHINESE_DIGITS.get(tens, 1) * 10 + CHINESE_DIGITS.get(ones, 0))
    elif len(digits) == 1:
        value = Decimal(CHINESE_DIGITS[digits])
    else:
        return None
    # 「一顆半」：半在單位字之後
    trailing_half = text[match.end():match.end() + 2]
    half = match.group("half") or (len(trailing_half) == 2 and trailing_half[1] == "半")
    return value + (Decimal("0.5") if half else 0)


def _quantity_text(raw_text: str, canonical_name: str | None) -> str:
    """只留下原文中的數量部分：去掉 [分組]、(備註) 與食材名稱，避免把「九層塔」「六份」當成數量。"""
    text = NOTE_RE.sub(" ", raw_text)
    if canonical_name:
        text = text.replace(canonical_name, " ")
    return text.strip()


def _normalize_quantity(record: dict[str, Any]) -> dict[str, Any]:
    raw_text = _text(record.get("raw_text")) or ""
    source_unit = _text(record.get("unit"))
    unit, unit_status, issues = _normalize_unit(source_unit, raw_text)
    source_amount = _lenient_decimal(record.get("amount"), "amount", issues)
    source_weight = _lenient_decimal(record.get("weight_g"), "weight_g", issues)

    contains_qualitative = any(word in raw_text for word in QUALITATIVE_WORDS)
    range_match = RANGE_RE.search(raw_text)
    fraction_match = FRACTION_RE.search(raw_text)
    has_digit = NUMBER_RE.search(raw_text) is not None
    amount: Decimal | None = None
    amount_min: Decimal | None = None
    amount_max: Decimal | None = None

    if contains_qualitative and not has_digit:
        if source_amount is not None:
            # 「適量／少許」等用詞：依資料整理時填入的 amount 作為參考用量（非原文明示）
            amount = source_amount
            quantity_status = "reference_amount"
            issues.append("qualitative_quantity_uses_reference_amount")
        else:
            quantity_status = "qualitative_unresolved"
            issues.append("qualitative_quantity_without_number")
    elif range_match:
        amount_min = Decimal(range_match.group("low"))
        amount_max = Decimal(range_match.group("high"))
        if amount_min > amount_max:
            amount_min, amount_max = amount_max, amount_min
        quantity_status = "source_range"
        issues.append("range_requires_usage_policy")
    elif fraction_match:
        denominator = Decimal(fraction_match.group("den"))
        if denominator == 0:
            quantity_status = "invalid"
            issues.append("invalid_fraction")
        else:
            amount = Decimal(fraction_match.group("num")) / denominator
            quantity_status = "source_exact"
    elif has_digit and source_amount is not None:
        amount = source_amount
        quantity_status = "source_exact"
    else:
        quantity_text = _quantity_text(raw_text, _text(record.get("canonical_name")))
        chinese = _chinese_number(quantity_text)
        vague = VAGUE_QUANTITY_RE.search(quantity_text) is not None
        if chinese is not None and not vague and source_amount == chinese:
            # 原文中文數字與 amount 一致才算原文明示
            amount = source_amount
            quantity_status = "source_exact"
        elif source_amount is not None and (chinese is None or vague):
            # 原文沒寫數量或只是概略用量：依資料整理時填入的 amount 作為參考用量
            amount = source_amount
            quantity_status = "reference_amount"
            issues.append("unstated_quantity_uses_reference_amount" if chinese is None
                          else "vague_quantity_uses_reference_amount")
        elif source_amount is not None:
            # 以 Excel 的 amount 為準（工作簿已人工確認）；原文數字不同只留提醒
            amount = source_amount
            quantity_status = "reference_amount"
            issues.append("raw_text_differs_from_amount")
        else:
            # 原文與 amount 都沒有可用數量
            quantity_status = "unresolved"
            issues.append("quantity_not_supported_by_raw_text")

    weight_g: Decimal | None = None
    if unit == "g" and amount is not None and quantity_status == "source_exact":
        weight_g = amount
        weight_status = "source_explicit"
    elif unit == "g" and amount is not None and quantity_status == "reference_amount":
        weight_g = amount
        weight_status = "reference_amount"
    elif unit == "g" and quantity_status == "source_range":
        weight_status = "source_range"
    elif source_weight is not None:
        weight_status = "unverified_conversion"
        issues.append("source_weight_requires_conversion_evidence")
    else:
        weight_status = "unresolved"

    return {
        "source_amount": source_amount,
        "source_unit": source_unit,
        "source_weight_g": source_weight,
        "amount": amount,
        "amount_min": amount_min,
        "amount_max": amount_max,
        "unit": unit,
        "quantity_status": quantity_status,
        "unit_status": unit_status,
        "weight_g": weight_g,
        "weight_status": weight_status,
        "quality_issues": sorted(set(issues)),
    }


def _parse_datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=timezone.utc)
    text_value = str(value).strip()
    try:
        parsed = datetime.fromisoformat(text_value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise WorkbookValidationError(f"Invalid datetime value: {value!r}") from exc
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def _method_candidates(instructions: Iterable[str]) -> list[str]:
    text_value = "\n".join(instructions)
    return [method for method in METHOD_KEYWORDS if method in text_value]


def _ingredient_similarity(left: set[str], right: set[str]) -> float:
    union = left | right
    if not union:
        return 0.0
    return len(left & right) / len(union)


def _suspected_duplicates(
    recipe_rows: list[dict[str, Any]], ingredients_by_recipe: dict[str, set[str]]
) -> list[dict[str, Any]]:
    pairs: list[dict[str, Any]] = []
    for index, left in enumerate(recipe_rows):
        for right in recipe_rows[index + 1 :]:
            if left["standard_dish_id"] == right["standard_dish_id"]:
                continue
            left_name = left["normalized_name"]
            right_name = right["normalized_name"]
            name_score = SequenceMatcher(None, left_name, right_name).ratio()
            ingredient_score = _ingredient_similarity(
                ingredients_by_recipe.get(left["recipe_id"], set()),
                ingredients_by_recipe.get(right["recipe_id"], set()),
            )
            contains = (
                min(len(left_name), len(right_name)) >= 3
                and (left_name in right_name or right_name in left_name)
            )
            if (name_score >= 0.72 and ingredient_score >= 0.35) or (
                contains and ingredient_score >= 0.35
            ):
                pairs.append(
                    {
                        "recipe_id_a": left["recipe_id"],
                        "recipe_id_b": right["recipe_id"],
                        "name_similarity": round(name_score, 4),
                        "ingredient_jaccard": round(ingredient_score, 4),
                        "status": "manual_review_required",
                    }
                )
    return pairs


def _catalog_records(workbook: Any) -> list[dict[str, Any]]:
    if CATALOG_SHEET not in workbook.sheetnames:
        return []
    rows = workbook[CATALOG_SHEET].iter_rows(values_only=True)
    headers = [str(h) if h is not None else "" for h in next(rows, ())]
    missing = CATALOG_COLUMNS - set(headers)
    if missing:
        raise WorkbookValidationError(f"Worksheet {CATALOG_SHEET} missing columns: {', '.join(sorted(missing))}")
    records = []
    for values in rows:
        record = dict(zip(headers, values))
        name, food_id = _text(record.get("合併標準食材名稱")), _text(record.get("Food_ID"))
        if name is None and food_id is None:
            continue
        if not (name and food_id):
            raise WorkbookValidationError(f"{CATALOG_SHEET} has a row with a blank name or Food_ID")
        records.append({"food_id": food_id, "name": name})
    return records


def _apply_source_rights(source_input: list[dict[str, Any]], source_rights: dict[str, dict] | None) -> None:
    """依來源設定補上授權資訊（工作簿本身未附授權欄位時）。工作簿已填的值優先。"""
    if not source_rights:
        return
    fields = ("license_name", "rights_status", "rights_uri", "source_version", "retrieved_at")
    for row in source_input:
        defaults = source_rights.get(_text(row.get("source_name")) or "")
        if not defaults:
            continue
        for field in fields:
            if _text(row.get(field)) is None and defaults.get(field) is not None:
                row[field] = defaults[field]


REVIEW_FIELDS = {
    # 審核紀錄的 approvals 鍵 → (工作表, 欄位)
    "source_review": ("source_archive", "review_status"),
    "recipe_review": ("recipes", "review_status"),
    "cooking_method": ("recipes", "cooking_method_status"),
    "step_review": ("recipe_steps", "review_status"),
    "adjustable_range": ("recipe_ingredients", "adjustment_status"),
}


def load_review_record(path: str | Path) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _apply_review_record(raw_sheets: dict[str, list[dict[str, Any]]], record: dict[str, Any] | None,
                         source_sha256: str) -> dict[str, Any] | None:
    """套用人工審核紀錄（工作簿沒有審核欄位時）。

    紀錄綁定工作簿 SHA-256：工作簿內容一改，舊的審核不會套用到新內容。
    工作簿已填的審核欄位優先。
    """
    if not record:
        return None
    summary = {key: record.get(key) for key in ("reviewed_on", "reviewer", "scope", "workbook_sha256")}
    if record.get("workbook_sha256") != source_sha256:
        return {**summary, "applied": False, "reason": "workbook SHA-256 differs from the reviewed version"}
    for key, status in (record.get("approvals") or {}).items():
        sheet, field = REVIEW_FIELDS[key]
        for row in raw_sheets[sheet]:
            if _text(row.get(field)) is None:
                row[field] = status
    return {**summary, "applied": True, "approvals": record.get("approvals"),
            "distinct_recipe_pairs": record.get("distinct_recipe_pairs") or []}


def _reviewed_distinct_pairs(review_applied: dict[str, Any] | None) -> set[frozenset[str]]:
    """人工確認「不是重複」的疑似重複配對；只在審核紀錄符合工作簿時生效。"""
    if not (review_applied or {}).get("applied"):
        return set()
    pairs = set()
    for pair in review_applied["distinct_recipe_pairs"]:
        ids = frozenset(_text(recipe_id) for recipe_id in pair)
        if len(pair) != 2 or len(ids) != 2 or None in ids:
            raise WorkbookValidationError(f"Invalid distinct recipe pair in review record: {pair!r}")
        pairs.add(ids)
    return pairs


def load_source_rights(path: str | Path) -> dict[str, dict]:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    return {name: entry for name, entry in data["sources"].items()}


def parse_recipe_workbook(path: str | Path, source_rights: dict[str, dict] | None = None,
                          review_record: dict[str, Any] | None = None) -> ImportBundle:
    workbook_path = Path(path).resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)

    source_sha256 = _file_hash(workbook_path)
    import_batch_id = source_sha256
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        raw_sheets = {
            sheet_name: _sheet_records(workbook, sheet_name)
            for sheet_name in REQUIRED_COLUMNS
        }
        catalog_input = _catalog_records(workbook)
    finally:
        workbook.close()

    review_applied = _apply_review_record(raw_sheets, review_record, source_sha256)
    source_input = raw_sheets["source_archive"]
    _apply_source_rights(source_input, source_rights)
    recipe_input = raw_sheets["recipes"]
    step_input = raw_sheets["recipe_steps"]
    ingredient_input = raw_sheets["recipe_ingredients"]
    alias_input = raw_sheets["ingredient_aliases"]
    allergen_input = raw_sheets["allergens"]

    def require_unique(records: list[dict[str, Any]], key: str, sheet: str) -> None:
        values = [_text(record.get(key)) for record in records]
        missing = [record["__row__"] for record, value in zip(records, values) if value is None]
        duplicates = sorted(value for value, count in Counter(values).items() if value and count > 1)
        if missing:
            raise WorkbookValidationError(f"{sheet}.{key} blank at rows: {missing}")
        if duplicates:
            raise WorkbookValidationError(f"{sheet}.{key} duplicates: {duplicates[:10]}")

    require_unique(source_input, "raw_id", "source_archive")
    require_unique(recipe_input, "recipe_id", "recipes")
    require_unique(step_input, "step_id", "recipe_steps")
    require_unique(ingredient_input, "recipe_ingredient_id", "recipe_ingredients")
    require_unique(allergen_input, "allergen_id", "allergens")

    source_by_raw_id = {_text(row["raw_id"]): row for row in source_input}
    recipe_ids = {_text(row["recipe_id"]) for row in recipe_input}
    orphan_raw_ids = sorted(
        _text(row["raw_id"])
        for row in recipe_input
        if _text(row["raw_id"]) not in source_by_raw_id
    )
    orphan_step_ids = sorted(
        _text(row["recipe_id"])
        for row in step_input
        if _text(row["recipe_id"]) not in recipe_ids
    )
    orphan_ingredient_ids = sorted(
        _text(row["recipe_id"])
        for row in ingredient_input
        if _text(row["recipe_id"]) not in recipe_ids
    )
    if orphan_raw_ids or orphan_step_ids or orphan_ingredient_ids:
        raise WorkbookValidationError(
            "Workbook contains orphan references: "
            f"raw={orphan_raw_ids[:10]}, steps={orphan_step_ids[:10]}, "
            f"ingredients={orphan_ingredient_ids[:10]}"
        )

    source_rows: list[dict[str, Any]] = []
    source_quality_by_raw: dict[str, list[str]] = {}
    for row in source_input:
        raw_id = _text(row["raw_id"])
        assert raw_id is not None
        rights_status = (_text(row.get("rights_status")) or "unknown").lower()
        review_status = (_text(row.get("review_status")) or "pending").lower()
        issues: list[str] = []
        if not _text(row.get("license_name")):
            issues.append("missing_license")
        if rights_status not in {"verified", "approved"}:
            issues.append("rights_not_verified")
        if not _text(row.get("source_version")):
            issues.append("missing_source_version")
        if not _parse_datetime(row.get("retrieved_at")):
            issues.append("missing_retrieved_at")
        if review_status != "approved":
            issues.append("source_review_pending")
        source_quality_by_raw[raw_id] = issues
        content_fields = {
            key: row.get(key)
            for key in (
                "source_name",
                "source_url",
                "source_recipe_id",
                "recipe_name_raw",
                "servings_raw",
                "ingredients_raw",
                "steps_raw",
                "tags_raw",
            )
        }
        source_rows.append(
            {
                "raw_id": raw_id,
                "source_name": _text(row["source_name"]),
                "source_url": _text(row["source_url"]),
                "source_recipe_id": _text(row["source_recipe_id"]),
                "recipe_name_raw": _text(row["recipe_name_raw"]),
                "servings_raw": _text(row["servings_raw"]),
                "ingredients_raw": _text(row["ingredients_raw"]),
                "steps_raw": _text(row["steps_raw"]),
                "tags_raw": _text(row["tags_raw"]),
                "license_name": _text(row.get("license_name")),
                "rights_uri": _text(row.get("rights_uri")),
                "rights_status": rights_status,
                "retrieved_at": _parse_datetime(row.get("retrieved_at")),
                "source_version": _text(row.get("source_version")),
                "content_hash": _content_hash(content_fields),
                "import_batch_id": import_batch_id,
                "review_status": review_status,
                "quality_issues": issues,
            }
        )

    step_rows: list[dict[str, Any]] = []
    instructions_by_recipe: dict[str, list[str]] = defaultdict(list)
    step_numbers: dict[str, set[int]] = defaultdict(set)
    for row in step_input:
        recipe_id = _text(row["recipe_id"])
        instruction = _text(row["instruction_clean"])
        if recipe_id is None or instruction is None:
            raise WorkbookValidationError(f"Invalid recipe step at row {row['__row__']}")
        step_no = _integer(row["step_no"], "step_no")
        if step_no in step_numbers[recipe_id]:
            raise WorkbookValidationError(f"Duplicate step number {recipe_id}/{step_no}")
        step_numbers[recipe_id].add(step_no)
        instructions_by_recipe[recipe_id].append(instruction)
        cooking_method = _text(row.get("cooking_method"))
        step_rows.append(
            {
                "step_id": _text(row["step_id"]),
                "recipe_id": recipe_id,
                "step_no": step_no,
                "instruction_clean": instruction,
                "cooking_method": cooking_method,
                "method_status": (_text(row.get("method_status")) or "unreviewed").lower(),
                "review_status": (_text(row.get("review_status")) or "pending").lower(),
            }
        )

    ingredient_rows: list[dict[str, Any]] = []
    ingredients_by_recipe: dict[str, set[str]] = defaultdict(set)
    ingredient_numbers: dict[str, set[Decimal]] = defaultdict(set)
    food_names: dict[str, set[str]] = defaultdict(set)
    ingredient_issues_by_recipe: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in ingredient_input:
        ingredient_id = _text(row["recipe_ingredient_id"])
        recipe_id = _text(row["recipe_id"])
        food_id = _text(row["Food_ID"])
        raw_text = _text(row["raw_text"])
        canonical_name = _text(row["canonical_name"])
        if not all((ingredient_id, recipe_id, food_id, raw_text, canonical_name)):
            raise WorkbookValidationError(f"Invalid recipe ingredient at row {row['__row__']}")
        sequence = _decimal(row["sequence"])
        if sequence is None or sequence <= 0:
            raise WorkbookValidationError(
                f"sequence must be a positive number: {row['sequence']!r}"
            )
        if sequence in ingredient_numbers[recipe_id]:
            raise WorkbookValidationError(f"Duplicate ingredient sequence {recipe_id}/{sequence}")
        ingredient_numbers[recipe_id].add(sequence)
        ingredients_by_recipe[recipe_id].add(food_id)
        food_names[food_id].add(canonical_name)
        quantity = _normalize_quantity(row)
        adjustment_status = (_text(row.get("adjustment_status")) or "not_reviewed").lower()
        adjustable_min = _lenient_decimal(row.get("adjustable_min"), "adjustable_min", quantity["quality_issues"])
        adjustable_max = _lenient_decimal(row.get("adjustable_max"), "adjustable_max", quantity["quality_issues"])
        if adjustable_min is not None and adjustable_max is not None and adjustable_min > adjustable_max:
            quantity["quality_issues"].append("invalid_adjustable_range")
            adjustable_min = adjustable_max = None
        source_nutrition = {
            f"source_{field}": _lenient_decimal(row.get(field), field, quantity["quality_issues"])
            for field in ("protein_g", "carbohydrates_g", "fat_g", "sugar_g", "sodium_mg")
        }
        if adjustment_status != "approved":
            quantity["quality_issues"].append("adjustable_range_not_reviewed")
        quality_issues = sorted(set(quantity.pop("quality_issues")))
        if quality_issues:
            ingredient_issues_by_recipe[recipe_id].append(
                {
                    "recipe_ingredient_id": ingredient_id,
                    "food_id": food_id,
                    "raw_text": raw_text,
                    "issues": quality_issues,
                }
            )
        ingredient_rows.append(
            {
                "recipe_ingredient_id": ingredient_id,
                "recipe_id": recipe_id,
                "sequence": sequence,
                "raw_text": raw_text,
                "canonical_name": canonical_name,
                "food_id": food_id,
                **quantity,
                "adjustable_min": adjustable_min,
                "adjustable_max": adjustable_max,
                "adjustment_status": adjustment_status,
                "is_required": row.get("is_required") if isinstance(row.get("is_required"), bool) else None,
                "source_protein_g": source_nutrition["source_protein_g"],
                "source_carbohydrates_g": source_nutrition["source_carbohydrates_g"],
                "source_fat_g": source_nutrition["source_fat_g"],
                "source_sugar_g": source_nutrition["source_sugar_g"],
                "source_sodium_mg": source_nutrition["source_sodium_mg"],
                "allergens": _split_list(row.get("allergens")),
                "quality_issues": quality_issues,
            }
        )

    # 標準食材字典：食譜沒用到的食材也收進 ingredient_aliases，作為全系統共用的 FOOD 字典
    catalog_ids = set()
    for entry in catalog_input:
        if entry["food_id"] in catalog_ids:
            raise WorkbookValidationError(f"{CATALOG_SHEET} duplicate Food_ID: {entry['food_id']}")
        catalog_ids.add(entry["food_id"])
        food_names[entry["food_id"]].add(entry["name"])
    if catalog_ids:
        missing_from_catalog = sorted(set(food_names) - catalog_ids)
        if missing_from_catalog:
            raise WorkbookValidationError(
                f"Food_ID used in recipes but missing from {CATALOG_SHEET}: {missing_from_catalog[:10]}")

    canonical_conflicts = {
        food_id: sorted(names) for food_id, names in food_names.items() if len(names) > 1
    }
    if canonical_conflicts:
        raise WorkbookValidationError(
            f"Food_ID canonical name conflicts: {dict(list(canonical_conflicts.items())[:10])}"
        )

    alias_targets: dict[str, set[str]] = defaultdict(set)
    aliases_by_food: dict[str, set[str]] = defaultdict(set)
    for row in alias_input:
        food_id = _text(row["Food_ID"])
        alias = _text(row["alias"])
        canonical_name = _text(row["canonical_name"])
        if not all((food_id, alias, canonical_name)):
            raise WorkbookValidationError(f"Invalid ingredient alias at row {row['__row__']}")
        if food_id in food_names and canonical_name not in food_names[food_id]:
            raise WorkbookValidationError(
                f"Ingredient alias canonical name conflicts for {food_id}: {canonical_name}"
            )
        alias_targets[normalize_name(alias)].add(food_id)
        aliases_by_food[food_id].add(alias)

    ingredient_alias_rows: list[dict[str, Any]] = []
    ambiguous_alias_count = 0
    for food_id in sorted(food_names):
        aliases: list[str] = []
        ambiguous_aliases: list[str] = []
        for alias in sorted(aliases_by_food.get(food_id, set())):
            if len(alias_targets[normalize_name(alias)]) == 1:
                aliases.append(alias)
            else:
                ambiguous_aliases.append(alias)
                ambiguous_alias_count += 1
        issues = ["ambiguous_aliases_require_review"] if ambiguous_aliases else []
        ingredient_alias_rows.append(
            {
                "food_id": food_id,
                "canonical_name": next(iter(food_names[food_id])),
                "aliases": aliases,
                "ambiguous_aliases": ambiguous_aliases,
                "review_status": "pending",
                "quality_issues": issues,
            }
        )

    recipe_rows: list[dict[str, Any]] = []
    recipe_by_id: dict[str, dict[str, Any]] = {}
    normalized_groups: dict[str, list[str]] = defaultdict(list)
    for row in recipe_input:
        recipe_id = _text(row["recipe_id"])
        raw_id = _text(row["raw_id"])
        recipe_name = _text(row["recipe_name"])
        if not all((recipe_id, raw_id, recipe_name)):
            raise WorkbookValidationError(f"Invalid recipe at row {row['__row__']}")
        source = source_by_raw_id[raw_id]
        servings, servings_status, servings_issues = _parse_servings(row["servings_raw"])
        normalized = normalize_name(recipe_name)
        normalized_groups[normalized].append(recipe_id)
        raw_name = _text(source["recipe_name_raw"])
        name_aliases = [raw_name] if raw_name and raw_name != recipe_name else []
        supplied_methods = _split_list(row.get("cooking_methods"))
        methods = supplied_methods or _method_candidates(instructions_by_recipe.get(recipe_id, []))
        method_status = (_text(row.get("cooking_method_status")) or "unreviewed").lower()
        review_status = (_text(row.get("review_status")) or "pending").lower()
        quality_issues = list(source_quality_by_raw[raw_id]) + servings_issues
        if not instructions_by_recipe.get(recipe_id):
            quality_issues.append("missing_steps")
        if not ingredients_by_recipe.get(recipe_id):
            quality_issues.append("missing_ingredients")
        if not methods:
            quality_issues.append("missing_cooking_method")
        if method_status != "approved":
            quality_issues.append("cooking_method_not_reviewed")
        if review_status != "approved":
            quality_issues.append("recipe_review_pending")
        for ingredient_issue in ingredient_issues_by_recipe.get(recipe_id, []):
            for issue in ingredient_issue["issues"]:
                quality_issues.append(f"ingredient:{issue}")

        recipe_row = {
            "recipe_id": recipe_id,
            "raw_id": raw_id,
            "standard_dish_id": standard_dish_id(recipe_name),
            "recipe_name": recipe_name,
            "normalized_name": normalized,
            "name_aliases": name_aliases,
            "servings_raw": _text(row["servings_raw"]),
            "original_servings": servings,
            "servings_status": servings_status,
            "category": _text(row["category"]),
            "meal_type": _text(row["meal_type"]),
            "cuisine": _text(row["Cuisine"]),
            "source_name": _text(row["source_name"]),
            "source_url": _text(row["source_url"]),
            "allergens": _split_list(row.get("allergens")),
            "cooking_methods": methods,
            "cooking_method_status": method_status,
            "duplicate_status": "unique_candidate",
            "duplicate_of_recipe_id": None,
            "duplicate_candidate_recipe_ids": [],
            "review_status": review_status,
            "quality_status": "not_qualified",
            "is_eligible": False,
            "quality_issues": sorted(set(quality_issues)),
        }
        recipe_rows.append(recipe_row)
        recipe_by_id[recipe_id] = recipe_row

    input_recipe_by_id = {_text(row["recipe_id"]): row for row in recipe_input}
    exact_duplicate_groups: list[dict[str, Any]] = []
    for normalized_name, ids in sorted(normalized_groups.items()):
        if len(ids) <= 1:
            continue
        ids = sorted(ids)
        primary_id = ids[0]
        exact_duplicate_groups.append(
            {
                "recipe_ids": ids,
                "primary_recipe_id": primary_id,
                "status": "manual_review_required",
                "records": [
                    {
                        "recipe_id": recipe_id,
                        "raw_id": recipe_by_id[recipe_id]["raw_id"],
                        "recipe_name": recipe_by_id[recipe_id]["recipe_name"],
                        "recipes_row": input_recipe_by_id[recipe_id]["__row__"],
                        "source_archive_row": source_by_raw_id[recipe_by_id[recipe_id]["raw_id"]]["__row__"],
                        "source_url": recipe_by_id[recipe_id]["source_url"],
                    }
                    for recipe_id in ids
                ],
            }
        )
        for recipe_id in ids:
            recipe_by_id[recipe_id]["duplicate_status"] = "exact_name_duplicate"
            recipe_by_id[recipe_id]["duplicate_candidate_recipe_ids"] = [
                candidate for candidate in ids if candidate != recipe_id
            ]
            recipe_by_id[recipe_id]["quality_issues"].append("exact_name_duplicate_review_required")
            if recipe_id != primary_id:
                recipe_by_id[recipe_id]["duplicate_of_recipe_id"] = primary_id
                # 同一道菜只保留 primary 可計入合格，其餘不能拿來補足 150 道
                recipe_by_id[recipe_id]["quality_issues"].append("duplicate_not_primary")

    distinct_pairs = _reviewed_distinct_pairs(review_applied)
    suspected_pairs = _suspected_duplicates(recipe_rows, ingredients_by_recipe)
    for pair in suspected_pairs:
        if frozenset((pair["recipe_id_a"], pair["recipe_id_b"])) in distinct_pairs:
            pair["status"] = "reviewed_distinct"
            continue
        left = recipe_by_id[pair["recipe_id_a"]]
        right = recipe_by_id[pair["recipe_id_b"]]
        if left["duplicate_status"] == "unique_candidate":
            left["duplicate_status"] = "suspected_duplicate"
        if right["duplicate_status"] == "unique_candidate":
            right["duplicate_status"] = "suspected_duplicate"
        left["duplicate_candidate_recipe_ids"].append(right["recipe_id"])
        right["duplicate_candidate_recipe_ids"].append(left["recipe_id"])
        left["quality_issues"].append("suspected_duplicate_review_required")
        right["quality_issues"].append("suspected_duplicate_review_required")
        # 疑似重複在人工確認前只讓編號較小的一筆可計入合格
        later = max(left, right, key=lambda recipe: recipe["recipe_id"])
        later["quality_issues"].append("suspected_duplicate_not_primary")

    for recipe in recipe_rows:
        recipe["duplicate_candidate_recipe_ids"] = sorted(
            set(recipe["duplicate_candidate_recipe_ids"])
        )
        recipe["quality_issues"] = sorted(set(recipe["quality_issues"]))
        blocking = {
            "missing_license",
            "rights_not_verified",
            "missing_source_version",
            "missing_retrieved_at",
            "source_review_pending",
            "recipe_review_pending",
            "missing_steps",
            "missing_ingredients",
            "missing_cooking_method",
            "cooking_method_not_reviewed",
            "ingredient:invalid_unit",
            "ingredient:quantity_not_supported_by_raw_text",
            "ingredient:qualitative_quantity_without_number",
            "ingredient:adjustable_range_not_reviewed",
            "ingredient:invalid_adjustable_range",
            "nonpositive_servings",
            "duplicate_not_primary",
            "suspected_duplicate_not_primary",
        }
        recipe["is_eligible"] = not any(
            issue in blocking or issue.startswith("ingredient:invalid_numeric_value:")
            for issue in recipe["quality_issues"]
        )
        recipe["quality_status"] = "qualified" if recipe["is_eligible"] else "not_qualified"

    allergen_rows = []
    for row in allergen_input:
        allergen_id = _text(row["allergen_id"])
        allergen_name = _text(row["allergen_name"])
        reference = _text(row["reference"])
        if not all((allergen_id, allergen_name, reference)):
            raise WorkbookValidationError(f"Invalid allergen at row {row['__row__']}")
        allergen_rows.append(
            {
                "allergen_id": allergen_id,
                "allergen_name": allergen_name,
                "reference": reference,
                "source_version": _text(row.get("source_version")),
                "review_status": (_text(row.get("review_status")) or "pending").lower(),
            }
        )

    qualified = sum(1 for recipe in recipe_rows if recipe["is_eligible"])
    reason_counts = Counter(
        issue for recipe in recipe_rows if not recipe["is_eligible"] for issue in recipe["quality_issues"]
    )
    exact_duplicate_records = sum(len(group["recipe_ids"]) - 1 for group in exact_duplicate_groups)
    ingredient_ids_by_reason: dict[str, list[str]] = defaultdict(list)
    for recipe_issues in ingredient_issues_by_recipe.values():
        for issue in recipe_issues:
            for reason in issue["issues"]:
                ingredient_ids_by_reason[reason].append(issue["recipe_ingredient_id"])
    rights_reasons = {"missing_license", "rights_not_verified", "missing_source_version", "missing_retrieved_at"}
    remaining_blockers = sorted(
        {issue for recipe in recipe_rows if not recipe["is_eligible"] for issue in recipe["quality_issues"]}
        & (rights_reasons | {"source_review_pending", "recipe_review_pending", "cooking_method_not_reviewed",
                             "ingredient:adjustable_range_not_reviewed"})
    )
    blocker_text = (
        "Remaining blockers: " + ", ".join(remaining_blockers) + "; unverified records are not counted as qualified."
        if remaining_blockers else "Quality issues remain; unverified records are not counted as qualified."
    )

    report = {
        "schema_version": SCHEMA_VERSION,
        "source_file": workbook_path.name,
        "source_sha256": source_sha256,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "counts": {
            "acquired_recipes": len(recipe_rows),
            "deduplicated_standard_dishes": len(recipe_rows) - exact_duplicate_records,
            "exact_duplicate_records": exact_duplicate_records,
            "exact_duplicate_groups": len(exact_duplicate_groups),
            "suspected_duplicate_pairs": len(suspected_pairs),
            "qualified_recipes": qualified,
            "not_qualified_recipes": len(recipe_rows) - qualified,
            "recipe_steps": len(step_rows),
            "recipe_ingredients": len(ingredient_rows),
            "distinct_food_ids": len(ingredient_alias_rows),
            "catalog_food_ids": len(catalog_ids),
            "ambiguous_alias_assignments": ambiguous_alias_count,
        },
        "not_qualified_reason_counts": dict(sorted(reason_counts.items())),
        "duplicate_review": {
            "exact_name_groups": exact_duplicate_groups,
            "suspected_pairs": suspected_pairs,
        },
        "manual_review": {
            "source_records": [
                {"raw_id": row["raw_id"], "issues": row["quality_issues"]}
                for row in source_rows
                if row["quality_issues"]
            ],
            "recipes": [
                {
                    "recipe_id": row["recipe_id"],
                    "issues": row["quality_issues"],
                }
                for row in recipe_rows
                if row["quality_issues"]
            ],
            "ingredient_ids_by_reason": {
                reason: sorted(set(ingredient_ids))
                for reason, ingredient_ids in sorted(ingredient_ids_by_reason.items())
            },
            "food_aliases": [
                {
                    "food_id": row["food_id"],
                    "ambiguous_alias_count": len(row["ambiguous_aliases"]),
                }
                for row in ingredient_alias_rows
                if row["ambiguous_aliases"]
            ],
        },
        "review_record": review_applied,
        "g01": {
            "status": "passed" if qualified >= 150 else "not_passed",
            "required_qualified_recipes": 150,
            "actual_qualified_recipes": qualified,
            "gap": max(0, 150 - qualified),
            "reason": (
                "At least 150 recipes have complete provenance and review evidence."
                if qualified >= 150
                else blocker_text
            ),
        },
        "ac05": {
            "status": "passed" if qualified >= 150 else "not_passed",
            "reason": (
                "Import evidence includes provenance, quantities, steps, deduplication and qualified counts."
                if qualified >= 150
                else "The import is reproducible, but " + blocker_text
            ),
        },
    }

    return ImportBundle(
        source_archive=source_rows,
        recipes=recipe_rows,
        recipe_steps=step_rows,
        recipe_ingredients=ingredient_rows,
        ingredient_aliases=ingredient_alias_rows,
        allergens=allergen_rows,
        report=report,
    )


def select_unique_recipe_bundle(bundle: ImportBundle) -> ImportBundle:
    """Import one source for each exact normalized name; keep the lowest recipe ID.

    Full parsing and nutrition provenance still use the unfiltered workbook. Only
    the recipe import view is filtered, with the source rows recorded for audit.
    Similar names are not treated as confirmed duplicates.
    """
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in bundle.recipes:
        groups[normalize_name(row["recipe_name"])].append(row)
    skipped_ids = {row["recipe_id"] for rows in groups.values()
                   for row in sorted(rows, key=lambda r: r["recipe_id"])[1:]}
    if not skipped_ids and "import_selection" in bundle.report:
        return bundle
    locations = {
        record["recipe_id"]: record
        for group in bundle.report.get("duplicate_review", {}).get("exact_name_groups", [])
        for record in group.get("records", [])
    }

    def describe(row):
        return {key: row[key] for key in ("recipe_id", "raw_id", "recipe_name", "source_url")} | locations.get(row["recipe_id"], {})

    selections = []
    for name, rows in sorted(groups.items()):
        if len(rows) > 1:
            ordered = sorted(rows, key=lambda row: row["recipe_id"])
            selections.append({"normalized_name": name, "kept": describe(ordered[0]),
                               "skipped": [describe(row) for row in ordered[1:]]})
    selected = []
    for row in bundle.recipes:
        if row["recipe_id"] in skipped_ids:
            continue
        values = {**row, "duplicate_candidate_recipe_ids": [
            key for key in row["duplicate_candidate_recipe_ids"] if key not in skipped_ids]}
        if len(groups[normalize_name(row["recipe_name"])]) > 1:
            values["duplicate_status"] = "selected_exact_name_duplicate"
            values["quality_issues"] = [issue for issue in row["quality_issues"]
                                        if issue != "exact_name_duplicate_review_required"]
        selected.append(values)
    selected_ids = {row["recipe_id"] for row in selected}
    skipped_raw_ids = {row["raw_id"] for row in bundle.recipes if row["recipe_id"] in skipped_ids}
    sources = [row for row in bundle.source_archive if row["raw_id"] not in skipped_raw_ids]
    steps = [row for row in bundle.recipe_steps if row["recipe_id"] in selected_ids]
    ingredients = [row for row in bundle.recipe_ingredients if row["recipe_id"] in selected_ids]
    selection = {
        "policy": "exact_normalized_name_keep_lowest_recipe_id",
        "input_recipes": len(bundle.recipes), "imported_recipes": len(selected),
        "skipped_duplicate_recipes": len(skipped_ids), "imported_sources": len(sources),
        "imported_steps": len(steps), "imported_ingredients": len(ingredients),
        "imported_qualified_recipes": sum(row["is_eligible"] for row in selected),
        "imported_not_qualified_recipes": sum(not row["is_eligible"] for row in selected),
        "selections": selections,
    }
    return replace(bundle, recipes=selected, source_archive=sources, recipe_steps=steps,
                   recipe_ingredients=ingredients, report={**bundle.report, "import_selection": selection})


def _upsert_rows(
    session: Session, model: type[Any], rows: Iterable[dict[str, Any]], primary_key: str
) -> dict[str, int]:
    inserted = 0
    updated = 0
    for values in rows:
        identity = values[primary_key]
        instance = session.get(model, identity)
        if instance is None:
            session.add(model(**values))
            inserted += 1
        else:
            for key, value in values.items():
                setattr(instance, key, value)
            updated += 1
    return {"inserted": inserted, "updated": updated}


def _replace_children(
    session: Session, model: type[Any], rows: list[dict[str, Any]], primary_key: str,
    recipe_ids: set[str],
) -> dict[str, int]:
    """以工作簿為準整批取代這些食譜的子資料（步驟／食材）。

    只做 upsert 會留下工作簿已移除的列，步驟重新編號時也會撞到
    (recipe_id, step_no) 唯一限制；因此先刪除再寫入，仍在同一個交易內。
    """
    key_column = getattr(model, primary_key)
    existing = set(session.scalars(
        select(key_column).where(model.recipe_id.in_(recipe_ids))
    )) if recipe_ids else set()
    if recipe_ids:
        session.execute(delete(model).where(model.recipe_id.in_(recipe_ids)))
        session.flush()
    new_ids = {row[primary_key] for row in rows}
    session.add_all(model(**row) for row in rows)
    return {
        "inserted": len(new_ids - existing),
        "updated": len(new_ids & existing),
        "deleted": len(existing - new_ids),
    }


def import_recipe_bundle(session: Session, bundle: ImportBundle) -> dict[str, dict[str, int]]:
    """Idempotently import a parsed bundle into the six T01 tables.

    The caller owns the transaction. A failed import can therefore be rolled back as
    one unit, and no source workbook is ever changed. Steps and ingredients of every
    recipe in the workbook are replaced as a whole, so rows removed from the workbook
    do not linger in the database.
    """

    bundle = select_unique_recipe_bundle(bundle)
    results: dict[str, dict[str, int]] = {}
    # 工作簿已刪除的整道食譜：連同步驟、食材一併刪除（子表為 ON DELETE CASCADE），
    # 指向它的 duplicate_of 會變成 NULL；工作簿是食譜庫的唯一來源。
    recipe_ids = {row["recipe_id"] for row in bundle.recipes}
    removed_recipes = list(session.scalars(select(Recipe.recipe_id).where(Recipe.recipe_id.not_in(recipe_ids))))
    cascaded = {"recipe_steps": 0, "recipe_ingredients": 0}
    if removed_recipes:
        for key, model in (("recipe_steps", RecipeStep), ("recipe_ingredients", RecipeIngredient)):
            cascaded[key] = session.scalar(
                select(func.count()).select_from(model).where(model.recipe_id.in_(removed_recipes)))
        session.execute(delete(Recipe).where(Recipe.recipe_id.in_(removed_recipes)))
        session.flush()
    results["source_archive"] = _upsert_rows(
        session, SourceArchive, bundle.source_archive, "raw_id"
    )
    results["ingredient_aliases"] = _upsert_rows(
        session, IngredientAlias, bundle.ingredient_aliases, "food_id"
    )
    results["allergens"] = _upsert_rows(session, Allergen, bundle.allergens, "allergen_id")
    session.flush()
    # duplicate_of_recipe_id 指向同一張表：先全部寫入、再補上參照，
    # 不依賴工作簿列的順序（避免 primary 排在後面時違反外鍵）
    results["recipes"] = _upsert_rows(
        session, Recipe,
        [{**row, "duplicate_of_recipe_id": None} for row in bundle.recipes], "recipe_id",
    )
    session.flush()
    for row in bundle.recipes:
        if row.get("duplicate_of_recipe_id"):
            session.get(Recipe, row["recipe_id"]).duplicate_of_recipe_id = row["duplicate_of_recipe_id"]
    session.flush()
    results["recipes"]["deleted"] = len(removed_recipes)
    # 不再被任何食譜使用、也不在工作簿中的來源紀錄
    raw_ids = {row["raw_id"] for row in bundle.source_archive}
    removed_sources = session.execute(
        delete(SourceArchive)
        .where(SourceArchive.raw_id.not_in(raw_ids))
        .where(~SourceArchive.raw_id.in_(select(Recipe.raw_id)))
    ).rowcount
    results["source_archive"]["deleted"] = removed_sources
    results["recipe_steps"] = _replace_children(
        session, RecipeStep, list(bundle.recipe_steps), "step_id", recipe_ids
    )
    results["recipe_ingredients"] = _replace_children(
        session, RecipeIngredient, list(bundle.recipe_ingredients), "recipe_ingredient_id", recipe_ids
    )
    # 隨整道食譜一起刪除的子資料也計入 deleted
    for key, count in cascaded.items():
        results[key]["deleted"] += count
    session.flush()
    return results


def write_quality_report(report: dict[str, Any], path: str | Path) -> Path:
    report_path = Path(path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = report_path.with_suffix(report_path.suffix + ".tmp")
    temporary_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    temporary_path.replace(report_path)
    return report_path

