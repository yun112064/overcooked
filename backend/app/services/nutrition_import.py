"""Read the three workbooks and import them atomically without merging variants."""
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import insert, select
from sqlalchemy.orm import Session

from app.models.nutrition import FoodNutrition, NutritionImport, NutritionSource, NutritionSourceLink
from app.models.recipe import IngredientAlias, Recipe
from app.services.recipe_import import (ImportBundle, WorkbookValidationError, import_recipe_bundle,
                                       parse_recipe_workbook, select_unique_recipe_bundle)

SCHEMA_VERSION = "food-nutrition-import-v1"
NUTRIENT_COLUMNS = {
    "熱量(kcal/100g)": "energy_kcal",
    "蛋白質(g/100g)": "protein_g",
    "脂肪(g/100g)": "fat_g",
    "碳水化合物(g/100g)": "carbohydrate_g",
    "糖(g/100g)": "sugars_g",
    "鈉(mg/100g)": "sodium_mg",
    "膳食纖維(g/100g)": "dietary_fiber_g",
}


@dataclass(frozen=True)
class NutritionBundle:
    recipe_bundle: ImportBundle
    import_id: str
    source_files: dict
    observations: list[dict]
    sources: list[dict]
    links: list[dict]
    report: dict


def _text(value):
    return str(value).strip() if value is not None and str(value).strip() else None


def _hash(*values):
    return hashlib.sha256(json.dumps(values, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()


def _rows(path, sheet, required):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in book.sheetnames:
            raise WorkbookValidationError(f"Missing worksheet: {sheet}")
        rows = book[sheet].iter_rows(values_only=True)
        headers = next(rows, ())
        nonblank = [h for h in headers if h is not None]
        if len(set(nonblank)) != len(nonblank):
            raise WorkbookValidationError(f"Duplicate headers in {sheet}")
        missing = set(required) - set(headers)
        if missing:
            raise WorkbookValidationError(f"{sheet} missing columns: {sorted(missing)}")
        result = []
        for number, values in enumerate(rows, 2):
            row = {h: v for h, v in zip(headers, values) if h is not None}
            if any(v is not None for v in row.values()):
                result.append((number, row))
        if not result:
            raise WorkbookValidationError(f"Empty worksheet: {sheet}")
        return result
    finally:
        book.close()


def _number(value, location):
    if value is None or value == "":
        return None
    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise WorkbookValidationError(f"Invalid number at {location}: {value!r}") from exc
    if not number.is_finite() or number < 0 or number >= Decimal("1e12"):
        raise WorkbookValidationError(f"Out-of-range number at {location}: {value!r}")
    return number


def _raw(row):
    return json.loads(json.dumps(row, ensure_ascii=False, default=str, allow_nan=False))


def _references(value):
    # These workbooks explicitly separate source keys with newlines.
    refs = [key.strip() for key in str(value or "").splitlines() if key.strip()]
    if len(set(refs)) != len(refs):
        raise WorkbookValidationError("Duplicate source reference within a nutrition row")
    return refs


def parse_nutrition_workbooks(recipes_path, catalog_path, nutrition_path, *, source_rights=None, review_record=None):
    paths = {"recipes": Path(recipes_path), "catalog": Path(catalog_path), "nutrition": Path(nutrition_path)}
    files = {kind: {"name": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
             for kind, path in paths.items()}
    import_id = _hash(SCHEMA_VERSION, {kind: item["sha256"] for kind, item in files.items()})
    recipes = parse_recipe_workbook(paths["recipes"], source_rights, review_record)
    catalog = {}
    for number, row in _rows(paths["catalog"], "標準食材整合", {"Food_ID", "合併標準食材名稱"}):
        food_id, name = _text(row["Food_ID"]), _text(row["合併標準食材名稱"])
        if not food_id or not name or len(food_id) > 32 or food_id in catalog:
            raise WorkbookValidationError(f"Invalid or duplicate catalog entry at row {number}")
        catalog[food_id] = name
    recipe_catalog = {row["food_id"]: row["canonical_name"] for row in recipes.ingredient_aliases}
    if catalog != recipe_catalog:
        raise WorkbookValidationError("External catalog and recipe workbook food dictionary differ")

    required = set(NUTRIENT_COLUMNS) | {"Food_ID", "標準食材名稱", "熟製狀態", "烹飪方式", "食用前處理",
                                          "計算方式", "原來源ID", "計算來源ID", "原來源筆數", "計算來源筆數"}
    all_rows = _rows(paths["nutrition"], "All", required)
    source_rows = _rows(paths["nutrition"], "原始來源與判定",
                        set(NUTRIENT_COLUMNS) | {"原來源ID", "資料來源", "來源資料ID", "來源網址"})
    recipe_source_rows = _rows(paths["recipes"], "recipe_ingredients", {"recipe_ingredient_id", "Food_ID"})
    recipe_sources = {_text(row["recipe_ingredient_id"]): (number, row) for number, row in recipe_source_rows}
    sources = {}
    for number, row in source_rows:
        key, provider, record_id = (_text(row[field]) for field in ("原來源ID", "資料來源", "來源資料ID"))
        if not key or len(key) > 100 or key in sources or provider not in {"TFDA", "USDA"} or not record_id:
            raise WorkbookValidationError(f"Invalid or duplicate nutrition source at row {number}")
        for column in NUTRIENT_COLUMNS:
            _number(row[column], f"原始來源與判定/{number}/{column}")
        sources[key] = {
            "source_id": _hash(import_id, "source", key), "import_id": import_id,
            "source_key": key, "provider": provider, "source_record_id": record_id,
            "source_url": _text(row["來源網址"]), "source_sheet": "原始來源與判定",
            "source_row": number, "raw_data": _raw(row),
        }

    observations, links, name_differences, count_mismatches = [], [], [], []
    variant_keys = set()
    for number, row in all_rows:
        food_id, name, state = (_text(row[field]) for field in ("Food_ID", "標準食材名稱", "熟製狀態"))
        method = _text(row["烹飪方式"])
        if food_id not in catalog or not name or not state:
            raise WorkbookValidationError(f"Unknown food ID or blank name/state at All/{number}: {food_id}")
        variant = (food_id, state, method)
        if variant in variant_keys:
            raise WorkbookValidationError(f"Duplicate nutrition variant at All/{number}: {variant}")
        variant_keys.add(variant)
        values = {field: _number(row[column], f"All/{number}/{column}") for column, field in NUTRIENT_COLUMNS.items()}
        issues = [f"missing:{field}" for field, value in values.items() if value is None]
        if catalog[food_id] != name:
            issues.append("catalog_name_difference")
            name_differences.append({"food_id": food_id, "catalog_name": catalog[food_id],
                                     "nutrition_name": name, "source_row": number})
        original, calculated = _references(row["原來源ID"]), set(_references(row["計算來源ID"]))
        if not original or not calculated.issubset(original):
            raise WorkbookValidationError(f"Invalid source references at All/{number}")
        for column, expected in (("原來源筆數", len(original)), ("計算來源筆數", len(calculated))):
            declared = _number(row[column], f"All/{number}/{column}")
            if declared is None or declared != declared.to_integral_value():
                raise WorkbookValidationError(f"Invalid source count at All/{number}/{column}")
            if declared != expected:
                issues.append(f"source_count_mismatch:{column}")
                count_mismatches.append({"source_row": number, "food_id": food_id, "column": column,
                                         "declared": int(declared), "actual": expected})
        nutrition_id = _hash(import_id, "All", number)
        for key in original:
            if key in recipe_sources:
                source_number, source = recipe_sources[key]
                if _text(source["Food_ID"]) != food_id:
                    raise WorkbookValidationError(f"Recipe source food differs at All/{number}: {key}")
                issues.append("recipe_derived_nutrition_unreviewed")
                if key not in sources:
                    sources[key] = {
                        "source_id": _hash(import_id, "source", key), "import_id": import_id,
                        "source_key": key, "provider": "recipe_dataset", "source_record_id": key,
                        "source_url": None, "source_sheet": "recipe_ingredients", "source_row": source_number,
                        "raw_data": _raw(source),
                    }
            if key not in sources:
                raise WorkbookValidationError(f"Unknown source reference at All/{number}: {key}")
            links.append({"nutrition_id": nutrition_id, "source_id": sources[key]["source_id"],
                          "included_in_calculation": key in calculated})
        observations.append({
            "nutrition_id": nutrition_id, "import_id": import_id, "food_id": food_id, "source_row": number,
            "standard_name": name, "cooking_state": state, "cooking_method": method,
            "preparation": _text(row["食用前處理"]), "calculation_method": _text(row["計算方式"]),
            **values, "nutrition_complete": all(value is not None for value in values.values()),
            "quality_issues": sorted(set(issues)), "raw_data": _raw(row),
        })

    food_counts = Counter(row["food_id"] for row in observations)
    used = {row["food_id"] for row in recipes.recipe_ingredients}
    missing_ids = sorted(catalog.keys() - food_counts.keys())
    selected_recipes = select_unique_recipe_bundle(recipes)
    report = {
        "schema_version": SCHEMA_VERSION, "import_id": import_id, "source_files": files,
        "counts": {"catalog_foods": len(catalog), "nutrition_records": len(observations),
                   "nutrition_foods": len(food_counts), "foods_with_variants": sum(n > 1 for n in food_counts.values()),
                   "original_sources": len(source_rows), "recipe_sources": len(sources) - len(source_rows),
                   "source_links": len(links), "complete_nutrition_records": sum(r["nutrition_complete"] for r in observations),
                   "catalog_ids_without_nutrition": len(missing_ids), "name_difference_records": len(name_differences),
                   "recipe_foods": len(used), "recipe_foods_without_nutrition": len(used - food_counts.keys()),
                   "source_count_mismatches": len(count_mismatches)},
        "missing_nutrients": {field: sum(r[field] is None for r in observations) for field in NUTRIENT_COLUMNS.values()},
        "reconciliation": {"catalog_ids_without_nutrition": [{"food_id": key, "name": catalog[key]} for key in missing_ids],
                           "name_differences": name_differences, "source_count_mismatches": count_mismatches},
        "recipe_quality": selected_recipes.report,
    }
    # Bind all reads to exactly the version identified in the report.
    if any(hashlib.sha256(path.read_bytes()).hexdigest() != files[kind]["sha256"] for kind, path in paths.items()):
        raise WorkbookValidationError("A source workbook changed during validation; retry the import")
    return NutritionBundle(selected_recipes, import_id, files, observations, list(sources.values()), links, report)


def import_nutrition_bundle(session: Session, bundle: NutritionBundle):
    """The caller owns the transaction; previous nutrition versions are never removed."""
    selected = select_unique_recipe_bundle(bundle.recipe_bundle)
    incoming = {row["recipe_id"] for row in selected.recipes}
    # Only duplicates explicitly selected out of this same workbook may be deleted.
    skipped = {row["recipe_id"] for group in selected.report["import_selection"]["selections"]
               for row in group["skipped"]}
    if set(session.scalars(select(Recipe.recipe_id))) - incoming - skipped:
        raise WorkbookValidationError("Import would remove existing recipes; refusing to delete them")
    catalog = {row["food_id"]: row["canonical_name"] for row in bundle.recipe_bundle.ingredient_aliases}
    for food_id, name in session.execute(select(IngredientAlias.food_id, IngredientAlias.canonical_name)):
        if food_id in catalog and catalog[food_id] != name:
            raise WorkbookValidationError(f"Existing food dictionary name differs: {food_id}")
    result = {"recipes": import_recipe_bundle(session, selected)}
    if session.get(NutritionImport, bundle.import_id) is not None:
        result["nutrition"] = {"status": "already_imported", "import_id": bundle.import_id}
        return result
    session.add(NutritionImport(import_id=bundle.import_id, source_files=bundle.source_files,
                               report=bundle.report, imported_at=datetime.now(timezone.utc)))
    session.flush()
    # Explicit NULLs must be preserved in bulk inserts with different missing fields.
    for model, rows in ((NutritionSource, bundle.sources), (FoodNutrition, bundle.observations), (NutritionSourceLink, bundle.links)):
        for offset in range(0, len(rows), 500):
            session.execute(insert(model).execution_options(render_nulls=True), rows[offset:offset + 500])
        session.flush()
    result["nutrition"] = {"status": "imported", "import_id": bundle.import_id,
                           "records": len(bundle.observations), "sources": len(bundle.sources), "links": len(bundle.links)}
    return result


def find_nutrition_candidates(session: Session, *, import_id: str, food_id: str,
                              cooking_state: str, cooking_method: str | None):
    """Require an explicit version and exact state/method; never choose another state."""
    return list(session.scalars(select(FoodNutrition).where(
        FoodNutrition.import_id == import_id, FoodNutrition.food_id == food_id,
        FoodNutrition.cooking_state == cooking_state, FoodNutrition.cooking_method == cooking_method,
    ).order_by(FoodNutrition.source_row)))
