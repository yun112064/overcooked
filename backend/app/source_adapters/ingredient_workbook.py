"""Build the T02 lookup scope from recipe Food_ID values and workbook names."""

from __future__ import annotations

import unicodedata
from collections import OrderedDict, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .price_models import IngredientQuery


RECIPE_SHEET_NAME = "recipe_ingredients"
ALIAS_SHEET_NAME = "ingredient_aliases"
RECIPE_REQUIRED_COLUMNS = {"canonical_name", "Food_ID"}
ALIAS_REQUIRED_COLUMNS = {"canonical_name", "Food_ID", "alias"}


def _clean(value: object) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def _read_rows(workbook, sheet_name: str, required: set[str]):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"required worksheet not found: {sheet_name}")
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [_clean(value) for value in next(rows)]
    except StopIteration as exc:
        raise ValueError(f"worksheet {sheet_name} is empty") from exc
    missing = required.difference(headers)
    if missing:
        raise ValueError(
            f"{sheet_name} missing required columns: "
            + ", ".join(sorted(missing))
        )
    return rows, {name: headers.index(name) for name in required}


def load_ingredient_queries(path: str | Path) -> list[IngredientQuery]:
    """Return one query per standard name referenced by recipe Food_ID values.

    ``recipe_ingredients.Food_ID`` defines the collection scope.  Names are
    resolved from the workbook's canonical-name columns, while
    ``ingredient_aliases.alias`` remains a matching hint only.  Multiple Food_ID
    values sharing one canonical name are grouped into a single retailer query.
    """

    workbook = load_workbook(
        filename=Path(path), read_only=True, data_only=True
    )
    try:
        recipe_rows, recipe_indexes = _read_rows(
            workbook, RECIPE_SHEET_NAME, RECIPE_REQUIRED_COLUMNS
        )
        food_id_names: OrderedDict[str, list[str]] = OrderedDict()
        for row in recipe_rows:
            food_id = _clean(row[recipe_indexes["Food_ID"]])
            if not food_id:
                continue
            names = food_id_names.setdefault(food_id, [])
            canonical = _clean(row[recipe_indexes["canonical_name"]])
            if canonical and canonical not in names:
                names.append(canonical)

        alias_rows, alias_indexes = _read_rows(
            workbook, ALIAS_SHEET_NAME, ALIAS_REQUIRED_COLUMNS
        )
        alias_names: dict[str, list[str]] = defaultdict(list)
        aliases: dict[str, list[str]] = defaultdict(list)
        for row in alias_rows:
            food_id = _clean(row[alias_indexes["Food_ID"]])
            if not food_id:
                continue
            canonical = _clean(row[alias_indexes["canonical_name"]])
            alias = _clean(row[alias_indexes["alias"]])
            if canonical and canonical not in alias_names[food_id]:
                alias_names[food_id].append(canonical)
            if alias and alias not in aliases[food_id]:
                aliases[food_id].append(alias)

        grouped: OrderedDict[str, dict[str, object]] = OrderedDict()
        for food_id, recipe_names in food_id_names.items():
            candidates: list[str] = []
            for name in [*recipe_names, *alias_names[food_id]]:
                if name and name not in candidates:
                    candidates.append(name)
            if len(candidates) != 1:
                reason = (
                    "Food_ID has no canonical_name in workbook"
                    if not candidates
                    else "Food_ID has conflicting canonical_name values: "
                    + " | ".join(candidates)
                )
                grouped[f"__unmapped__:{food_id}"] = {
                    "canonical_name": f"Food_ID:{food_id}",
                    "food_ids": [food_id],
                    "aliases": [],
                    "mapping_missing_reason": reason,
                }
                continue

            canonical = candidates[0]
            values = grouped.setdefault(
                canonical,
                {
                    "canonical_name": canonical,
                    "food_ids": [],
                    "aliases": [],
                    "mapping_missing_reason": None,
                },
            )
            food_ids = values["food_ids"]
            assert isinstance(food_ids, list)
            food_ids.append(food_id)
            match_aliases = values["aliases"]
            assert isinstance(match_aliases, list)
            for alias in aliases[food_id]:
                if alias != canonical and alias not in match_aliases:
                    match_aliases.append(alias)
    finally:
        workbook.close()

    return [IngredientQuery(**metadata) for metadata in grouped.values()]
