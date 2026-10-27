"""File adapters for the common knowledge record contract; no DB/model calls."""
import json
from pathlib import Path

from openpyxl import load_workbook

from knowledge_rag.store import validate_record

HEADERS = {
    '知識ID': 'knowledge_id', '食材ID': 'ingredient_id', '食材名稱': 'ingredient_name',
    '來源版本': 'source_version', '知識內容': 'fact_zh', '原文片段': 'excerpt',
    '來源網址': 'source_url', '使用條件': 'rights', '適用條件': 'conditions',
    '專業人士姓名': 'professional_name', '專業人士身分': 'professional_role',
    '文章標題': 'title', '原文定位': 'locator', '證據類型': 'evidence_type',
    '審核狀態': 'status',
}
EXCEL_FIELDS = set(HEADERS.values())


def _excel_records(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        names = [name for name in ('knowledge', '食材知識') if name in workbook.sheetnames]
        if len(names) == 1:
            sheet = workbook[names[0]]
        elif not names and len(workbook.sheetnames) == 1:
            sheet = workbook.active
        else:
            raise ValueError('Use one knowledge sheet named knowledge or 食材知識')
        if sheet is None:
            raise ValueError('Knowledge worksheet required')
        rows = sheet.iter_rows()
        header = next(rows, None)
        if header is None:
            raise ValueError('Knowledge headers required in row 1')
        columns = []
        for cell in header:
            if cell.data_type == 'f':
                raise ValueError('Knowledge headers cannot contain formulas')
            name = str(cell.value).strip() if cell.value is not None else ''
            field = HEADERS.get(name, name)
            if field and field not in EXCEL_FIELDS:
                raise ValueError(f'Unknown knowledge column: {name}')
            columns.append(field)
        populated = [field for field in columns if field]
        if len(populated) != len(set(populated)):
            raise ValueError('Duplicate knowledge columns')
        required = {'knowledge_id', 'ingredient_id', 'source_version', 'fact_zh',
                    'excerpt', 'source_url', 'rights', 'conditions'}
        if not required.issubset(populated):
            raise ValueError('Missing knowledge columns: ' + ', '.join(sorted(required - set(populated))))
        records = []
        for number, row in enumerate(rows, start=2):
            if all(cell.value is None for cell in row):
                continue
            record = {}
            for field, cell in zip(columns, row):
                if cell.data_type == 'f':
                    raise ValueError(f'Row {number}: formulas are not knowledge text')
                if cell.value is not None and not field:
                    raise ValueError(f'Row {number}: populated cell has no column header')
                if field and cell.value is not None:
                    if not isinstance(cell.value, str):
                        raise ValueError(f'Row {number}: {field} must be text')
                    record[field] = cell.value.strip()
            records.append(record)
        return records
    finally:
        workbook.close()


def load_records(path):
    """Read/validate a whole file before mutation; sheet status never approves it."""
    path = Path(path)
    if path.suffix.lower() == '.json':
        with path.open(encoding='utf-8-sig') as handle:
            document = json.load(handle)
        records = document.get('records') if isinstance(document, dict) else None
    elif path.suffix.lower() == '.xlsx':
        records = _excel_records(path)
    else:
        raise ValueError('Knowledge import supports .json and .xlsx')
    if not isinstance(records, list) or not records:
        raise ValueError('At least one knowledge record required')
    for number, record in enumerate(records, start=1):
        try:
            validate_record(record)
        except ValueError as exc:
            raise ValueError(f'Record {number}: {exc}') from None
    return records
